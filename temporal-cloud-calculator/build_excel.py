#!/usr/bin/env python3
"""Genera calculadora-temporal-cloud.xlsx (fórmulas nativas, celdas amarillas = input)."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

N_WF = 20
N_STEPS = 150

TIPOS = [
    ("Activity", 1, 2),
    ("Local Activities (1 por WFT)", 1, 1),
    ("Timer / timeout SDK", 1, 0),
    ("Signal", 1, 1),
    ("Signal externo", 1, 1),
    ("Query", 1, 0),
    ("Update", 1, 1),
    ("Child (hijo NO listado)", 2, 1),
    ("Continue-As-New", 1, 1),
    ("Upsert Search Attributes", 1, 1),
    ("Nexus Operation (caller)", 1, 1),
    ("WFT Heartbeat (LA largas)", 1, 0),
]

ORIGENES = ["Cliente", "Signal-With-Start", "Schedule", "Ninguno"]

# name, exec/mes, origen, Workflow Execution Duration (h), Payload StartWorkflow (KB)
SAMPLE_WF = [
    ("Pago interbancario", 50000, "Cliente", 0.4, 8),
    ("Alta de cliente", 8000, "Schedule", 2.5, 16),
]

# workflow, tipo, nombre, veces, %, retries, heartbeats, payload KB
SAMPLE_STEPS = [
    ("Pago interbancario", "Activity", "Validar cuenta", 1, 100, 0, 0, 4),
    ("Pago interbancario", "Activity", "Débito core", 1, 98, 0.3, 4, 8),
    ("Pago interbancario", "Activity", "Crédito destino", 1, 95, 0.4, 6, 8),
    ("Pago interbancario", "Activity", "Notificar", 1, 100, 0, 0, 2),
    ("Pago interbancario", "Timer / timeout SDK", "Espera SLA / AwaitWithTimeout", 1, 15, 0, 0, 0),
    ("Pago interbancario", "Signal", "Confirmación manual", 1, 8, 0, 0, 1),
    ("Pago interbancario", "Query", "Consulta estado UI", 2, 40, 0, 0, 2),
    ("Alta de cliente", "Activity", "KYC", 1, 100, 1, 12, 20),
    ("Alta de cliente", "Activity", "Crear producto", 1, 85, 0.2, 2, 8),
    ("Alta de cliente", "Timer / timeout SDK", "Espera documentos", 1, 30, 0, 0, 0),
    ("Alta de cliente", "Signal", "Documentos listos", 1, 30, 0, 0, 10),
    ("Alta de cliente", "Child (hijo NO listado)", "Onboarding cards (no listado)", 1, 60, 0, 0, 12),
]


def styles():
    thin = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )
    return {
        "title": Font(name="Calibri", size=16, bold=True, color="FFFFFF"),
        "h": Font(name="Calibri", size=10, bold=True, color="FFFFFF"),
        "cell": Font(name="Calibri", size=10, color="1D2939"),
        "mono": Font(name="Consolas", size=9, color="1D2939"),
        "thin": thin,
        "fill_title": PatternFill("solid", fgColor="1B4F72"),
        "fill_h": PatternFill("solid", fgColor="154360"),
        "fill_in": PatternFill("solid", fgColor="FFF3C4"),
        "fill_out": PatternFill("solid", fgColor="E8F6F3"),
        "fill_alt": PatternFill("solid", fgColor="F8FAFC"),
        "wrap": Alignment(wrap_text=True, vertical="center"),
        "right": Alignment(horizontal="right", vertical="center"),
    }


def header_row(ws, headers, s):
    for col, h in enumerate(headers, 1):
        c = ws.cell(1, col, h)
        c.font = s["h"]
        c.fill = s["fill_h"]
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = s["thin"]
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def widths(ws, wsizes):
    for i, w in enumerate(wsizes, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def paint(cell, s, fill, num=None):
    cell.font = s["cell"]
    cell.fill = fill
    cell.border = s["thin"]
    cell.alignment = s["wrap"]
    if num:
        cell.number_format = num


def sheet_tarifas(wb, s):
    ws = wb.active
    ws.title = "Tarifas"
    ws.sheet_properties.tabColor = "1B4F72"
    ws.merge_cells("A1:C1")
    t = ws["A1"]
    t.value = "Temporal Cloud — tarifas commitment 3 años (celdas amarillas = input)"
    t.font = s["title"]
    t.fill = s["fill_title"]
    ws.row_dimensions[1].height = 28

    ws["A3"] = "Parámetro"
    ws["B3"] = "Valor"
    ws["C3"] = "Nota"
    for col in range(1, 4):
        ws.cell(3, col).font = s["h"]
        ws.cell(3, col).fill = s["fill_h"]
        ws.cell(3, col).border = s["thin"]

    rows = [
        (4, "RATE_ACTIONS", "Actions ($ / millón)", 23.16, "Contrato. List público empieza en $50/millón."),
        (5, "RATE_ACTIVE", "Active Storage ($ / GB-h)", 0.039, "Contrato. List público $0.042/GBh."),
        (6, "RATE_RETAINED", "Retained Storage ($ / GB-h)", 0.00105, "Igual al list público."),
        (7, "HORAS_MES", "Horas / mes (referencia GB)", 730, "Temporal usa 744 h para GB incluidos del plan. No entra en Actions."),
        (8, "RETENTION_DIAS", "Namespace Retention Period (días) — solo Retained Storage", 30, "NO es Active. Tras close, el namespace guarda el Event History estos días (máx. 90). Un valor para todos los workflows."),
        (9, "PCT_CERT", "% certificación (sobre producción)", 15, "Atajo de volumen. USD cert = TOTAL prod × este % / 100. Mismo mix, Duration, Retention y retries que prod."),
        (10, "PCT_DEV", "% desarrollo (sobre producción)", 5, "Atajo de volumen. USD dev = TOTAL prod × este % / 100. TOTAL cuenta = prod + cert + dev."),
    ]
    for r, name, label, val, note in rows:
        ws.cell(r, 1, label)
        paint(ws.cell(r, 1), s, s["fill_alt"])
        c = ws.cell(r, 2, val)
        paint(c, s, s["fill_in"], "#,##0.00000")
        ws.cell(r, 3, note)
        paint(ws.cell(r, 3), s, s["fill_alt"])
        wb.defined_names.add(DefinedName(name, attr_text=f"Tarifas!$B${r}"))

    ws["B9"].number_format = "0.00"
    ws["B10"].number_format = "0.00"

    ws["A12"] = "Cómo leer"
    ws["A12"].font = Font(name="Calibri", size=12, bold=True, color="1B4F72")
    notes = [
        "1. El detalle de Workflows/Pasos es producción. Cert y dev se estiman como % del TOTAL prod (Tarifas B9/B10).",
        "2. Workflows: type, ejecuciones, origen, Workflow Execution Duration (h) = Active, Payload StartWorkflow (KB).",
        "3. Pasos: veces × (%/100) × unidad = Actions. Payload (KB) del mensaje que fluye en ese paso.",
        "4. Event History (KB) = Payload StartWorkflow + Σ (veces × % × Payload × copias). Activity ×2 (input+result) y ×(1+retries).",
        "5. Query y Timer no persisten Payload. Active = exec × GB × Duration. Retained = exec × GB × Retention × 24.",
        "6. Origen Cliente = 1 start. Signal-With-Start = 1. Schedule = 3 (incluye start). Child no listado = 2. Nombres unicos (SUMIF).",
        "7. Replay no factura. Heartbeat solo si llega al server (SDK ~80% del Heartbeat Timeout). 1 GB = 1.000.000 KB.",
        "8. Fuentes: https://docs.temporal.io/cloud/actions  y  https://docs.temporal.io/cloud/pricing",
    ]
    for i, line in enumerate(notes):
        ws.cell(13 + i, 1, line).alignment = s["wrap"]
        ws.merge_cells(start_row=13 + i, start_column=1, end_row=13 + i, end_column=3)
        ws.row_dimensions[13 + i].height = 28

    widths(ws, [48, 16, 88])
    ws.freeze_panes = "A4"
    return ws


def sheet_listas(wb, s):
    ws = wb.create_sheet("Listas")
    ws.sheet_properties.tabColor = "6941C6"
    ws["A1"] = "Origen"
    ws["B1"] = "Tipo de paso"
    ws["C1"] = "Actions unitarias"
    ws["D1"] = "Copias Payload en Event History"
    for col in range(1, 5):
        ws.cell(1, col).font = s["h"]
        ws.cell(1, col).fill = s["fill_h"]
        ws.cell(1, col).border = s["thin"]
    for i, o in enumerate(ORIGENES, 2):
        ws.cell(i, 1, o)
        paint(ws.cell(i, 1), s, s["fill_alt"])
    for i, (t, u, copies) in enumerate(TIPOS, 2):
        ws.cell(i, 2, t)
        ws.cell(i, 3, u)
        ws.cell(i, 4, copies)
        paint(ws.cell(i, 2), s, s["fill_alt"])
        paint(ws.cell(i, 3), s, s["fill_alt"], "0")
        paint(ws.cell(i, 4), s, s["fill_alt"], "0")
    wb.defined_names.add(DefinedName("ORIGENES", attr_text="Listas!$A$2:$A$5"))
    wb.defined_names.add(DefinedName("TIPOS", attr_text="Listas!$B$2:$B$13"))
    wb.defined_names.add(DefinedName("TIPO_UNIDAD", attr_text="Listas!$B$2:$C$13"))
    wb.defined_names.add(DefinedName("TIPO_COPIAS", attr_text="Listas!$B$2:$D$13"))
    widths(ws, [22, 36, 18, 32])
    return ws


def sheet_workflows(wb, s):
    ws = wb.create_sheet("Workflows")
    ws.sheet_properties.tabColor = "027A48"
    headers = [
        "Nombre (type)",
        "Ejecuciones / mes",
        "Origen",
        "Actions origen",
        "Workflow Execution Duration (h)",
        "Payload StartWorkflow (KB)",
        "Event History (KB)",
        "Actions/exec (pasos)",
        "Actions / mes",
        "Active Storage (GBh)",
        "Retained Storage (GBh)",
        "USD Actions",
        "USD Active",
        "USD Retained",
        "USD total",
    ]
    header_row(ws, headers, s)
    last = N_WF + 1
    for r in range(2, last + 1):
        i = r - 2
        sample = SAMPLE_WF[i] if i < len(SAMPLE_WF) else None
        name = sample[0] if sample else ""
        exec_n = sample[1] if sample else ""
        origin = sample[2] if sample else "Cliente"
        dur = sample[3] if sample else ""
        start_kb = sample[4] if sample else ""

        paint(ws.cell(r, 1, name), s, s["fill_in"])
        c = ws.cell(r, 2, exec_n if exec_n != "" else None)
        paint(c, s, s["fill_in"], "#,##0")
        paint(ws.cell(r, 3, origin if sample else ""), s, s["fill_in"])
        d = ws.cell(
            r,
            4,
            f'=IF(A{r}="","",IF(C{r}="Signal-With-Start",1,IF(C{r}="Schedule",3,IF(C{r}="Ninguno",0,1))))',
        )
        paint(d, s, s["fill_out"], "0")
        e = ws.cell(r, 5, dur if dur != "" else None)
        paint(e, s, s["fill_in"], "0.00")
        f = ws.cell(r, 6, start_kb if start_kb != "" else None)
        paint(f, s, s["fill_in"], "0.00")
        g = ws.cell(
            r,
            7,
            f'=IF(A{r}="","",MAX(0,N(F{r}))+SUMIF(Pasos!$A$2:$A${N_STEPS+1},A{r},Pasos!$K$2:$K${N_STEPS+1}))',
        )
        paint(g, s, s["fill_out"], "0.00")
        h = ws.cell(
            r,
            8,
            f'=IF(A{r}="","",SUMIF(Pasos!$A$2:$A${N_STEPS+1},A{r},Pasos!$J$2:$J${N_STEPS+1}))',
        )
        paint(h, s, s["fill_out"], "0.00")
        ii = ws.cell(r, 9, f'=IF(A{r}="","",MAX(0,N(B{r}))*(D{r}+H{r}))')
        paint(ii, s, s["fill_out"], "#,##0.00")
        j = ws.cell(r, 10, f'=IF(A{r}="","",MAX(0,N(B{r}))*(MAX(0,N(G{r}))/1000000)*MAX(0,N(E{r})))')
        paint(j, s, s["fill_out"], "#,##0.000")
        k = ws.cell(
            r,
            11,
            f'=IF(A{r}="","",MAX(0,N(B{r}))*(MAX(0,N(G{r}))/1000000)*MAX(0,N(RETENTION_DIAS))*24)',
        )
        paint(k, s, s["fill_out"], "#,##0.000")
        l = ws.cell(r, 12, f'=IF(A{r}="","",I{r}/1000000*RATE_ACTIONS)')
        paint(l, s, s["fill_out"], '"$"#,##0.00')
        m = ws.cell(r, 13, f'=IF(A{r}="","",J{r}*RATE_ACTIVE)')
        paint(m, s, s["fill_out"], '"$"#,##0.00')
        nn = ws.cell(r, 14, f'=IF(A{r}="","",K{r}*RATE_RETAINED)')
        paint(nn, s, s["fill_out"], '"$"#,##0.00')
        o = ws.cell(r, 15, f'=IF(A{r}="","",L{r}+M{r}+N{r})')
        paint(o, s, s["fill_out"], '"$"#,##0.00')
        ws.row_dimensions[r].height = 20

    wb.defined_names.add(DefinedName("WF_NAMES", attr_text=f"Workflows!$A$2:$A${last}"))
    dv = DataValidation(type="list", formula1="=Listas!$A$2:$A$5", allow_blank=True)
    dv.error = "Elige Cliente, Signal-With-Start, Schedule o Ninguno"
    dv.errorTitle = "Origen"
    dv.showErrorMessage = True
    dv.showDropDown = False
    ws.add_data_validation(dv)
    dv.add(f"C2:C{last}")
    ws.auto_filter.ref = f"A1:O{last}"
    widths(ws, [28, 18, 22, 14, 30, 26, 18, 18, 16, 20, 22, 14, 14, 14, 14])
    return ws


def sheet_pasos(wb, s):
    ws = wb.create_sheet("Pasos")
    ws.sheet_properties.tabColor = "FDB022"
    headers = [
        "Workflow (type)",
        "Tipo de paso",
        "Nombre del paso",
        "Veces si el camino ocurre",
        "% del flujo",
        "Retries avg (solo Activity)",
        "Heartbeats avg (solo Activity)",
        "Payload del mensaje (KB)",
        "Unidad oficial",
        "Actions / ejecución",
        "Event History (KB)",
    ]
    header_row(ws, headers, s)
    last = N_STEPS + 1
    for r in range(2, last + 1):
        i = r - 2
        sample = SAMPLE_STEPS[i] if i < len(SAMPLE_STEPS) else None
        if sample:
            ws.cell(r, 1, sample[0])
            ws.cell(r, 2, sample[1])
            ws.cell(r, 3, sample[2])
            ws.cell(r, 4, sample[3])
            ws.cell(r, 5, sample[4])
            ws.cell(r, 6, sample[5])
            ws.cell(r, 7, sample[6])
            ws.cell(r, 8, sample[7])
        else:
            for col in range(1, 9):
                ws.cell(r, col, None)
        for col in range(1, 9):
            paint(ws.cell(r, col), s, s["fill_in"])
        ws.cell(r, 4).number_format = "0.00"
        ws.cell(r, 5).number_format = "0"
        ws.cell(r, 6).number_format = "0.00"
        ws.cell(r, 7).number_format = "0.00"
        ws.cell(r, 8).number_format = "0.00"
        h = ws.cell(
            r,
            9,
            f'=IF(B{r}="","",IFERROR(VLOOKUP(B{r},Listas!$B$2:$C$13,2,0),1))',
        )
        paint(h, s, s["fill_out"], "0")
        ii = ws.cell(
            r,
            10,
            f'=IF(OR(A{r}="",B{r}=""),"",IF(B{r}="Activity",MAX(0,N(D{r}))*(MIN(100,MAX(0,N(E{r})))/100)*(1+MAX(0,N(F{r}))+MAX(0,N(G{r}))),MAX(0,N(D{r}))*(MIN(100,MAX(0,N(E{r})))/100)*I{r}))',
        )
        paint(ii, s, s["fill_out"], "0.000")
        kk = ws.cell(
            r,
            11,
            f'=IF(OR(A{r}="",B{r}=""),"",MAX(0,N(D{r}))*(MIN(100,MAX(0,N(E{r})))/100)*MAX(0,N(H{r}))*IFERROR(VLOOKUP(B{r},Listas!$B$2:$D$13,3,0),0)*IF(B{r}="Activity",1+MAX(0,N(F{r})),1))',
        )
        paint(kk, s, s["fill_out"], "0.000")
        ws.row_dimensions[r].height = 18

    dv_wf = DataValidation(type="list", formula1=f"=Workflows!$A$2:$A${N_WF+1}", allow_blank=True)
    dv_wf.error = "Elige un nombre de la hoja Workflows"
    dv_wf.errorTitle = "Workflow"
    dv_wf.showErrorMessage = True
    dv_wf.showDropDown = False
    ws.add_data_validation(dv_wf)
    dv_wf.add(f"A2:A{last}")
    dv_t = DataValidation(type="list", formula1="=Listas!$B$2:$B$13", allow_blank=True)
    dv_t.error = "Elige un tipo de la hoja Listas"
    dv_t.errorTitle = "Tipo"
    dv_t.showErrorMessage = True
    dv_t.showDropDown = False
    ws.add_data_validation(dv_t)
    dv_t.add(f"B2:B{last}")
    ws.auto_filter.ref = f"A1:K{last}"
    widths(ws, [26, 32, 36, 22, 14, 24, 26, 22, 16, 18, 18])
    return ws


def sheet_resumen(wb, s):
    ws = wb.create_sheet("Resumen", 1)  # after Tarifas; Workflows already exists
    ws.sheet_properties.tabColor = "B42318"
    ws.merge_cells("A1:D1")
    t = ws["A1"]
    t.value = "Resumen mensual — Temporal Cloud (commitment 3 años)"
    t.font = s["title"]
    t.fill = s["fill_title"]
    ws.row_dimensions[1].height = 28

    ws["A3"] = "Indicador"
    ws["B3"] = "Cantidad"
    ws["C3"] = "USD / mes"
    for col in range(1, 4):
        ws.cell(3, col).font = s["h"]
        ws.cell(3, col).fill = s["fill_h"]
        ws.cell(3, col).border = s["thin"]

    last_wf = N_WF + 1
    lines = [
        (4, "Tipos de workflow con nombre", f'=COUNTA(Workflows!A2:A{last_wf})', None),
        (5, "Ejecuciones / mes", f"=SUMIF(Workflows!A2:A{last_wf},\"<>\",Workflows!B2:B{last_wf})", None),
        (6, "Actions / mes", f"=SUMIF(Workflows!A2:A{last_wf},\"<>\",Workflows!I2:I{last_wf})", f"=B6/1000000*RATE_ACTIONS"),
        (7, "Actions (millones)", f"=B6/1000000", None),
        (8, "Actions / ejecución (media)", f'=IF(B5=0,0,B6/B5)', None),
        (9, "Active Storage (GBh)", f"=SUMIF(Workflows!A2:A{last_wf},\"<>\",Workflows!J2:J{last_wf})", f"=B9*RATE_ACTIVE"),
        (10, "Retained Storage (GBh)", f"=SUMIF(Workflows!A2:A{last_wf},\"<>\",Workflows!K2:K{last_wf})", f"=B10*RATE_RETAINED"),
        (11, "GB activo equivalente (GBh / horas mes)", f"=IF(HORAS_MES=0,0,B9/HORAS_MES)", None),
        (12, "TOTAL producción USD / mes", "-", f"=N(C6)+N(C9)+N(C10)"),
        (13, "Certificación (PCT_CERT % de producción)", f"=MAX(0,N(PCT_CERT))", f"=N(C12)*MAX(0,N(PCT_CERT))/100"),
        (14, "Desarrollo (PCT_DEV % de producción)", f"=MAX(0,N(PCT_DEV))", f"=N(C12)*MAX(0,N(PCT_DEV))/100"),
        (15, "TOTAL cuenta USD / mes (prod + cert + dev)", "-", f"=N(C12)+N(C13)+N(C14)"),
        (16, "Nombres duplicados en Workflows (0 = OK)", f'=SUMPRODUCT((Workflows!A2:A{last_wf}<>"")*(COUNTIF(Workflows!A2:A{last_wf},Workflows!A2:A{last_wf})>1))', None),
    ]
    for r, label, qty, usd in lines:
        ws.cell(r, 1, label)
        paint(ws.cell(r, 1), s, s["fill_alt"])
        paint(ws.cell(r, 2, qty), s, s["fill_out"], "#,##0.000")
        if usd:
            paint(ws.cell(r, 3, usd), s, s["fill_out"], '"$"#,##0.00')
        else:
            paint(ws.cell(r, 3, "-"), s, s["fill_alt"])

    ws["B5"].number_format = "#,##0"
    ws["B6"].number_format = "#,##0.00"
    ws["B4"].number_format = "0"
    ws["B13"].number_format = "0.00"
    ws["B14"].number_format = "0.00"
    ws["B16"].number_format = "0"
    for addr in ("A12", "B12", "C12"):
        ws[addr].fill = PatternFill("solid", fgColor="E8F6F3")
        ws[addr].font = Font(name="Calibri", size=11, bold=True)
    for addr in ("A15", "B15", "C15"):
        ws[addr].fill = PatternFill("solid", fgColor="AAF0C4")
        ws[addr].font = Font(name="Calibri", size=11, bold=True)

    ws["A18"] = "Costo por workflow type (producción): columnas L-O de la hoja Workflows."
    ws["A18"].font = Font(name="Calibri", size=12, bold=True, color="1B4F72")
    ws.merge_cells("A18:D18")

    ws["A20"] = (
        "Estimado de regimen, no la factura. Temporal documenta hasta 5% de diferencia entre metricas OpenMetrics y metering. "
        "Cert y dev son un % del TOTAL de produccion (mismo mix/Retention). No incluye Support, Export, Fairness ni TRU ociosos. "
        "Abrir en Excel para Windows o Microsoft 365: Habilitar edicion si aparece Vista protegida."
    )
    ws["A20"].alignment = s["wrap"]
    ws.merge_cells("A20:D20")
    ws.row_dimensions[20].height = 48
    widths(ws, [48, 18, 16, 16])
    return ws


def sheet_catalogo(wb, s):
    ws = wb.create_sheet("Catalogo")
    ws.sheet_properties.tabColor = "155EEF"
    ws.merge_cells("A1:D1")
    t = ws["A1"]
    t.value = "Catálogo oficial de Actions (referencia, no input)"
    t.font = s["title"]
    t.fill = s["fill_title"]
    headers = ["En la calculadora", "Usage / metric name", "Actions", "Nota oficial"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(3, col, h)
        c.font = s["h"]
        c.fill = s["fill_h"]
        c.border = s["thin"]
    rows = [
        ("Origen Cliente", "start_workflow", "1", "Start vía client. Dedup por Workflow ID no cuenta."),
        ("Origen Signal-With-Start", "signal_workflow_with_start", "1", "Un Action total, arranque o no."),
        ("Origen Schedule", "schedule start + start_workflow", "3", "Schedule Start cuenta 2; el workflow 1."),
        ("Activity", "schedule_activity + retry_activity + record_activity_heartbeat", "1 + retries + HB", "HB solo si llega al server. SDK throttle ~80% del timeout."),
        ("Local Activities (1 por WFT)", "record_activity_markers", "1 por Workflow Task", "N Local Activities del mismo WFT = 1 Action."),
        ("Timer / timeout SDK", "start_timer", "1", "Incluye AwaitWithTimeout / condition."),
        ("Signal", "signal_workflow", "1", "Client o workflow."),
        ("Signal externo", "signal_external_workflow", "1", "En el history del signaled."),
        ("Query", "query_workflow", "1", "Incluye call stack en UI Cloud. __temporal_workflow_metadata excluida."),
        ("Update", "accept_workflow_update / reject_workflow_update", "1", "Dedup por Update ID no cuenta."),
        ("Child (hijo NO listado)", "start_child_workflow + start", "2", "Intent + intento de start. Si el hijo es otra fila, no uses este tipo."),
        ("Continue-As-New", "continue_as_new_workflow", "1", "El nuevo run también tiene su start."),
        ("Upsert Search Attributes", "upsert_workflow_search_attributes", "1 por comando", "Varios SA en un comando = 1. SA al start no cuentan."),
        ("Nexus Operation (caller)", "schedule_nexus_operation", "1 en caller", "El handler factura sus primitives aparte."),
        ("WFT Heartbeat (LA largas)", "record_workflow_heartbeat", "1", "Local Activities largas que latean el WFT."),
        ("Replay", "-", "0", "No factura. Solo worker."),
    ]
    for i, row in enumerate(rows):
        r = 4 + i
        for col, val in enumerate(row, 1):
            ws.cell(r, col, val)
            paint(ws.cell(r, col), s, s["fill_alt"])
        ws.row_dimensions[r].height = 32
    ws["A21"] = "https://docs.temporal.io/cloud/actions"
    ws["A21"].hyperlink = "https://docs.temporal.io/cloud/actions"
    ws["A21"].font = Font(name="Calibri", size=10, color="155EEF", underline="single")
    ws["A22"] = "https://docs.temporal.io/cloud/pricing"
    ws["A22"].hyperlink = "https://docs.temporal.io/cloud/pricing"
    ws["A22"].font = Font(name="Calibri", size=10, color="155EEF", underline="single")
    widths(ws, [36, 52, 28, 72])
    return ws


def sheet_diccionario(wb, s):
    ws = wb.create_sheet("Diccionario")
    ws.sheet_properties.tabColor = "6941C6"
    ws.merge_cells("A1:G1")
    t = ws["A1"]
    t.value = "Diccionario del Excel — Temporal Cloud (commitment 3 años)"
    t.font = s["title"]
    t.fill = s["fill_title"]
    ws.row_dimensions[1].height = 28

    intro = [
        (
            "Event History",
            "Log append-only de una Workflow Execution: start, activities (input y result), signals, timers, etc. "
            "Es el dato que Temporal Cloud guarda en disco. NO se ingresa: se calcula de los Payloads. "
            "El mismo Event History (KB) se usa para Active y para Retained; lo que cambia es CUANTO TIEMPO se cobra.",
        ),
        (
            "Active Storage",
            "GB-hora del Event History MIENTRAS la ejecución está Running (open). "
            "Tiempo = Workflow Execution Duration (h). NO usa Namespace Retention Period.",
        ),
        (
            "Retained Storage",
            "GB-hora del Event History DESPUES de close, mientras el namespace lo retiene. "
            "Tiempo = Namespace Retention Period (días) × 24 h. NO usa Workflow Execution Duration.",
        ),
        (
            "Namespace Retention Period",
            "Setting del namespace (uno solo para todos los workflow types). Máx. 90 días en Temporal Cloud. "
            "Solo entra en Retained Storage. No hay un Retention de Active distinto: Active dura lo que dure el run.",
        ),
    ]
    ws["A3"] = "Conceptos (Active y Retained son dos cobros distintos)"
    ws["A3"].font = Font(name="Calibri", size=12, bold=True, color="1B4F72")
    ws.merge_cells("A3:G3")
    for i, (name, text) in enumerate(intro):
        r = 4 + i
        paint(ws.cell(r, 1, name), s, s["fill_alt"])
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        paint(ws.cell(r, 2, text), s, s["fill_alt"])
        ws.row_dimensions[r].height = 48

    headers = ["Término", "Hoja", "Columna / celda", "Tipo", "Aplica a", "Qué es", "Cómo se calcula / usa"]
    hr = 9
    for col, h in enumerate(headers, 1):
        c = ws.cell(hr, col, h)
        c.font = s["h"]
        c.fill = s["fill_h"]
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = s["thin"]
    ws.row_dimensions[hr].height = 24
    ws.freeze_panes = "A10"

    INP = "input (amarillo)"
    FOR = "fórmula (verde)"
    REF = "referencia"

    rows = [
        ("Actions ($ / millón)", "Tarifas", "B4  RATE_ACTIONS", INP, "Actions", "Precio del commitment por millón de Actions.", "Costo Actions = Actions/mes / 1e6 × esta tarifa."),
        ("Active Storage ($ / GB-h)", "Tarifas", "B5  RATE_ACTIVE", INP, "Active Storage", "Precio del GB-hora mientras el workflow está Running.", "Costo Active = Active GBh × esta tarifa."),
        ("Retained Storage ($ / GB-h)", "Tarifas", "B6  RATE_RETAINED", INP, "Retained Storage", "Precio del GB-hora del Event History ya cerrado.", "Costo Retained = Retained GBh × esta tarifa."),
        ("Horas / mes", "Tarifas", "B7  HORAS_MES", INP, "Ninguno (referencia)", "Convierte GBh a GB equivalente. Temporal usa 744 h para GB incluidos del plan.", "GB equiv. = Active GBh / Horas mes. No entra en Actions ni en el USD de storage."),
        ("Namespace Retention Period (días)", "Tarifas", "B8  RETENTION_DIAS", INP, "Solo Retained Storage", "Días que el namespace guarda el Event History tras close. Un valor para todos los types. Máx. 90.", "Retained GBh = exec × Event History GB × RETENTION_DIAS × 24. No entra en Active."),
        ("% certificación", "Tarifas", "B9  PCT_CERT", INP, "Costo cuenta", "Atajo de volumen sobre el TOTAL de producción.", "USD cert = C12 × PCT_CERT / 100. Asume mismo mix, Duration, Retention y retries."),
        ("% desarrollo", "Tarifas", "B10  PCT_DEV", INP, "Costo cuenta", "Atajo de volumen sobre el TOTAL de producción.", "USD dev = C12 × PCT_DEV / 100. TOTAL cuenta = prod + cert + dev."),
        ("Nombre (type)", "Workflows", "A", INP, "Todos", "Workflow Type. Debe ser único: Pasos se agrupa con SUMIF por este nombre.", "Si se duplica el nombre, los pasos se suman dos veces. Resumen B16 cuenta duplicados."),
        ("Ejecuciones / mes", "Workflows", "B", INP, "Actions + ambos storage", "Cuántas veces corre este type en el mes calendario.", "Multiplica Actions/exec, Active y Retained."),
        ("Origen", "Workflows", "C", INP, "Actions", "Cómo arranca: Cliente, Signal-With-Start, Schedule o Ninguno.", "Desplegable de Listas. Vacío o Cliente = 1 Action de start."),
        ("Actions origen", "Workflows", "D", FOR, "Actions", "Actions del arranque oficial.", "Cliente o Signal-With-Start = 1. Schedule = 3 (incluye start). Ninguno = 0."),
        ("Workflow Execution Duration (h)", "Workflows", "E", INP, "Solo Active Storage", "Horas que la ejecución está Running (open).", "Active GBh = exec × Event History GB × Duration. No entra en Retained."),
        ("Payload StartWorkflow (KB)", "Workflows", "F", INP, "Ambos storage", "Tamaño del input / mensaje de StartWorkflow.", "Se suma al Event History de cada ejecución."),
        ("Event History (KB)", "Workflows", "G", FOR, "Ambos storage", "Tamaño estimado del log persistido de UNA ejecución. No se escribe.", "F + SUMIF(Pasos Event History del mismo type). GB = KB / 1.000.000."),
        ("Actions/exec (pasos)", "Workflows", "H", FOR, "Actions", "Actions de los pasos, sin el start.", "SUMIF de Pasos J para este nombre."),
        ("Actions / mes", "Workflows", "I", FOR, "Actions", "Actions totales del type en el mes.", "Ejecuciones × (Actions origen + Actions/exec pasos)."),
        ("Active Storage (GBh)", "Workflows", "J", FOR, "Active Storage", "GB-hora mientras está open.", "exec × (Event History KB / 1e6) × Duration."),
        ("Retained Storage (GBh)", "Workflows", "K", FOR, "Retained Storage", "GB-hora tras close, régimen mensual.", "exec × (Event History KB / 1e6) × RETENTION_DIAS × 24."),
        ("USD Actions / Active / Retained / total", "Workflows", "L–O", FOR, "Costo", "Desglose USD del type.", "I/1e6×RATE_ACTIONS; J×RATE_ACTIVE; K×RATE_RETAINED; suma."),
        ("Workflow (type)", "Pasos", "A", INP, "Todos", "Debe coincidir exactamente con Workflows!A.", "Desplegable. SUMIF usa este texto."),
        ("Tipo de paso", "Pasos", "B", INP, "Actions + storage", "Activity, Timer, Signal, Query, etc.", "Desplegable. Define Actions unitarias y copias de Payload en Event History."),
        ("Nombre del paso", "Pasos", "C", INP, "Ninguno", "Etiqueta humana del paso.", "No entra en fórmulas."),
        ("Veces si el camino ocurre", "Pasos", "D", INP, "Actions + storage", "Cuántas veces se dispara el paso cuando el camino ocurre.", "Se multiplica por % y por unidad o Payload."),
        ("% del flujo", "Pasos", "E", INP, "Actions + storage", "Porcentaje de ejecuciones que toman ese camino. 0–100.", "veces × (min(100,%)/100)."),
        ("Retries avg (solo Activity)", "Pasos", "F", INP, "Actions + Event History", "Reintentos promedio de la Activity.", "Actions: +retries. Event History Activity: ×(1+retries). Ignorado en otros tipos."),
        ("Heartbeats avg (solo Activity)", "Pasos", "G", INP, "Solo Actions", "Heartbeats que llegan al server (SDK ~80% del Heartbeat Timeout).", "Suma Actions. No se suma al Event History (el payload del HB no se modela)."),
        ("Payload del mensaje (KB)", "Pasos", "H", INP, "Ambos storage", "Tamaño del mensaje que fluye en ese paso.", "Event History del paso = veces × % × Payload × copias × (1+retries si Activity)."),
        ("Unidad oficial", "Pasos", "I", FOR, "Actions", "Actions por ocurrencia según catálogo oficial.", "VLOOKUP del tipo en Listas C."),
        ("Actions / ejecución", "Pasos", "J", FOR, "Actions", "Actions esperadas de este paso por ejecución del type.", "Activity: veces×%×(1+retries+heartbeats). Otro: veces×%×unidad."),
        ("Event History (KB) del paso", "Pasos", "K", FOR, "Ambos storage", "Aporte de este paso al Event History de una ejecución.", "veces×%×Payload×copias (Listas D). Query y Timer copias=0. Activity ×2 y ×(1+retries)."),
        ("Origen / Tipo / copias", "Listas", "A–D", REF, "Lookups", "Tablas de los desplegables.", "C = Actions unitarias. D = copias del Payload en Event History (Activity=2, Query/Timer=0)."),
        ("Tipos de workflow con nombre", "Resumen", "B4", FOR, "Control", "Cuántas filas de Workflows tienen nombre.", "COUNTA(Workflows!A)."),
        ("Ejecuciones / mes", "Resumen", "B5", FOR, "Volumen", "Suma de ejecuciones con nombre.", "SUMIF Workflows B."),
        ("Actions / mes", "Resumen", "B6 / C6", FOR, "Actions", "Total Actions y su USD.", "SUMIF Workflows I. USD = B6/1e6×RATE_ACTIONS."),
        ("Active Storage (GBh)", "Resumen", "B9 / C9", FOR, "Active Storage", "Suma Active y su USD.", "SUMIF Workflows J × RATE_ACTIVE. Duration, no Retention."),
        ("Retained Storage (GBh)", "Resumen", "B10 / C10", FOR, "Retained Storage", "Suma Retained y su USD.", "SUMIF Workflows K × RATE_RETAINED. Retention, no Duration."),
        ("GB activo equivalente", "Resumen", "B11", FOR, "Referencia", "Active GBh pasado a GB con Horas/mes.", "B9 / HORAS_MES. No es un cobro extra."),
        ("TOTAL producción USD / mes", "Resumen", "C12", FOR, "Costo prod", "Actions + Active + Retained del detalle ingresado (producción).", "C6+C9+C10."),
        ("Certificación", "Resumen", "C13", FOR, "Costo cuenta", "% de producción.", "C12 × PCT_CERT / 100."),
        ("Desarrollo", "Resumen", "C14", FOR, "Costo cuenta", "% de producción.", "C12 × PCT_DEV / 100."),
        ("TOTAL cuenta USD / mes", "Resumen", "C15", FOR, "Costo", "Producción + certificación + desarrollo.", "C12+C13+C14. No incluye Support, Export, Fairness ni TRU ociosos."),
        ("Nombres duplicados", "Resumen", "B16", FOR, "Control", "0 = OK. Distinto de 0 = SUMIF doble cuenta Pasos.", "COUNTIF sobre Workflows A."),
        ("Catálogo Actions", "Catalogo", "A–D", REF, "Actions", "Nombres oficiales de usage. No es input.", "docs.temporal.io/cloud/actions"),
        ("Amarillo", "Todas", "—", INP, "—", "Celda editable.", "No pongas fórmulas en amarillo."),
        ("Verde", "Todas", "—", FOR, "—", "Celda calculada.", "No editar. Se recalcula al abrir (fullCalcOnLoad)."),
    ]

    for i, row in enumerate(rows):
        r = 10 + i
        for col, val in enumerate(row, 1):
            paint(ws.cell(r, col, val), s, s["fill_alt"] if i % 2 == 0 else PatternFill("solid", fgColor="FFFFFF"))
        ws.row_dimensions[r].height = 42

    last = 9 + len(rows)
    ws.auto_filter.ref = f"A9:G{last}"
    note_r = last + 2
    ws.cell(note_r, 1, "Fuentes: https://docs.temporal.io/cloud/actions  y  https://docs.temporal.io/cloud/pricing  ·  Estimado de régimen, no la factura (±5% metering). Ver tambien explicacion-excel.md")
    ws.merge_cells(start_row=note_r, start_column=1, end_row=note_r, end_column=7)
    ws.cell(note_r, 1).alignment = s["wrap"]
    widths(ws, [36, 14, 22, 18, 24, 52, 62])
    return ws


def main():
    wb = Workbook()
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, calcId=191029)
    wb.properties.creator = "ludhara-tech"
    wb.properties.title = "Calculadora Temporal Cloud"
    s = styles()
    sheet_tarifas(wb, s)
    sheet_listas(wb, s)
    sheet_workflows(wb, s)
    sheet_pasos(wb, s)
    sheet_resumen(wb, s)
    sheet_catalogo(wb, s)
    sheet_diccionario(wb, s)

    # defined name clash check
    banned = {"RF", "C", "R"}
    for n in list(wb.defined_names.keys()):
        if n in banned:
            raise SystemExit(f"defined name forbidden: {n}")

    out = Path(__file__).resolve().parent / "calculadora-temporal-cloud.xlsx"
    wb.save(out)
    print(f"wrote {out}: {N_WF} workflow slots, {N_STEPS} step rows")


if __name__ == "__main__":
    main()
