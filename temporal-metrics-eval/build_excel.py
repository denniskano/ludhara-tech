#!/usr/bin/env python3
"""Genera metricas-temporal-cloud-openmetrics.xlsx desde el catálogo oficial OpenMetrics v1."""

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

FUENTE = "https://docs.temporal.io/cloud/metrics/openmetrics/metrics-reference"
ENDPOINT = "https://metrics.temporal.io/v1/metrics"
DESCRIPTORES = "https://metrics.temporal.io/v1/descriptors"
CORTE = "2026-10-27"

COMMON_LABELS = "temporal_namespace, temporal_account, region"

# importancia: P1 paging / P2 degradación / P3 capacidad-FinOps / Dashboard / No alerta
# candidata: Sí solo P1–P3 (el propósito del archivo)

METRICS = [
    # --- Frontend ---
    dict(
        metrica="temporal_cloud_v1_service_request_count",
        categoria="Frontend Service",
        tipo="Rate",
        unidad="requests/s",
        desc_oficial="gRPC requests received per second.",
        desc_ops="Volumen de API. Denominador del error rate y del silencio de Start. Filtrar operation=StartWorkflowExecution|SignalWorkflowExecution|SignalWithStartWorkflowExecution para el SLO de ingreso.",
        labels="operation",
        optin="—",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2 (silencio Start) / denominador P1",
        condicion="Silencio: StartWorkflowExecution == 0 en horario hábil ≥15m. Como ratio: error_count/request_count > 1% con volumen > 0.1/s.",
        promql='sum by (temporal_namespace) (temporal_cloud_v1_service_request_count{operation="StartWorkflowExecution"}) == 0',
        duracion="15m",
        caveats="Rate ya promediada 1 min. Un burst sub-minuto no se ve en el valor medio. No usar rate().",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="Sí (StartSilence + denominador error rate)",
        dashboard="D01, D02",
        pareja="temporal_cloud_v1_service_error_count, temporal_cloud_v1_service_request_limit",
    ),
    dict(
        metrica="temporal_cloud_v1_service_request_throttled_count",
        categoria="Frontend Service",
        tipo="Rate",
        unidad="requests/s",
        desc_oficial="gRPC requests throttled per second. See Monitoring Trends Against Limits for guidance on setting alert thresholds against the corresponding limit metric.",
        desc_ops="Cloud está rechazando RPS del frontend. Paging: cualquier valor > 0 sostenido. El % medio vs limit puede seguir verde (burst absorbido en 1 min).",
        labels="operation",
        optin="—",
        importancia="P1",
        orden=1,
        candidata="Sí",
        severidad="P1",
        condicion="> 0 sostenido. Cruzar con service_request_count / service_request_limit.",
        promql="sum by (temporal_namespace) (temporal_cloud_v1_service_request_throttled_count) > 0",
        duracion="2m",
        caveats="Throttle puede ocurrir con utilización media < límite (agregación 1 min). No alertar solo el % vs limit.",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="Sí (TemporalCloudThrottled)",
        dashboard="D02",
        pareja="temporal_cloud_v1_service_request_limit, temporal_cloud_v1_service_request_count",
    ),
    dict(
        metrica="temporal_cloud_v1_service_error_count",
        categoria="Frontend Service",
        tipo="Rate",
        unidad="errors/s",
        desc_oficial="gRPC errors per second.",
        desc_ops="Errores de Frontend. Alertar como ratio vs request_count en Start/Signal, no el absoluto. Incluye errores de cliente y de servidor; filtrar operation.",
        labels="operation",
        optin="—",
        importancia="P1",
        orden=1,
        candidata="Sí",
        severidad="P1",
        condicion="error/request > 1% en Start|Signal|SignalWithStart y request > 0.1/s.",
        promql='(sum by (temporal_namespace) (temporal_cloud_v1_service_error_count{operation=~"StartWorkflowExecution|SignalWorkflowExecution|SignalWithStartWorkflowExecution"}) / clamp_min(sum by (temporal_namespace) (temporal_cloud_v1_service_request_count{operation=~"StartWorkflowExecution|SignalWorkflowExecution|SignalWithStartWorkflowExecution"}), 0.001)) > 0.01',
        duracion="10m",
        caveats="Sin umbral de volumen, namespaces idle disparan el ratio. No es health del worker.",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="Sí (TemporalCloudStartSignalErrorRate)",
        dashboard="D02",
        pareja="temporal_cloud_v1_service_request_count",
    ),
    dict(
        metrica="temporal_cloud_v1_service_pending_requests",
        categoria="Frontend Service",
        tipo="Value",
        unidad="pollers",
        desc_oficial="The number of pollers that are actively long polling for a task. Use this to track against temporal_cloud_v1_poller_limit.",
        desc_ops="Pollers en long-poll. Capacidad, no paging. Alertar % vs poller_limit.",
        labels="operation",
        optin="—",
        importancia="P3",
        orden=3,
        candidata="Sí",
        severidad="P3",
        condicion="pending / poller_limit > 80% ≥15m.",
        promql="sum by (temporal_namespace) (temporal_cloud_v1_service_pending_requests) / sum by (temporal_namespace) (temporal_cloud_v1_poller_limit) > 0.8",
        duracion="15m",
        caveats="Gauge. Un valor alto sin no_poller_tasks puede ser workers sobredimensionados, no saturación de matching.",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="Sí (TemporalCloudPollerLimitHigh)",
        dashboard="D01, D03",
        pareja="temporal_cloud_v1_poller_limit",
    ),
    dict(
        metrica="temporal_cloud_v1_resource_exhausted_error_count",
        categoria="Frontend Service",
        tipo="Rate",
        unidad="errors/s",
        desc_oficial="Resource exhaustion errors per second, incremented when a single resource receives a burst larger than it can absorb. SDKs retry these errors gracefully. This metric does not include throttling due to Namespace limits — see total_action_throttled_count and related throttle metrics.",
        desc_ops="Burst que Cloud no absorbe. Distinto de throttle de cuota APS/RPS. Paging aunque el % vs limit esté verde. Ver Detecting Resource Exhaustion.",
        labels="operation",
        optin="—",
        importancia="P1",
        orden=1,
        candidata="Sí",
        severidad="P1",
        condicion="> 0. SDK reintenta; si persiste hay hot key / partición / recurso interno saturado.",
        promql="sum by (temporal_namespace) (temporal_cloud_v1_resource_exhausted_error_count) > 0",
        duracion="2m",
        caveats="No es el mismo fenómeno que total_action_throttled_count. No silenciar porque 'el limit está holgado'.",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="Sí (TemporalCloudResourceExhausted)",
        dashboard="D02",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_service_latency_p50",
        categoria="Frontend Service",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 50th percentile latency of service requests in seconds.",
        desc_ops="Latencia Frontend mediana. Tablero. No alertar p50.",
        labels="operation",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta. Usar p95 con gate de request_count.",
        promql="—",
        duracion="—",
        caveats="No agregar percentiles entre dimensiones (operation, namespace, region): el p no es agregable. En ns de poco tráfico un solo request mueve p50/p95/p99.",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="No",
        dashboard="D01",
        pareja="temporal_cloud_v1_service_request_count",
    ),
    dict(
        metrica="temporal_cloud_v1_service_latency_p95",
        categoria="Frontend Service",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 95th percentile latency of service requests in seconds.",
        desc_ops="Candidata a alerta de latencia API solo con mínimo de samples (~20 req/min) y por operation, sin avg entre labels.",
        labels="operation",
        optin="—",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="p95 > umbral de SLO (p. ej. 0.5s) AND service_request_count de esa operation ≥ 20/min. No avg().",
        promql="temporal_cloud_v1_service_latency_p95{operation=\"StartWorkflowExecution\"} > 0.5 and temporal_cloud_v1_service_request_count{operation=\"StartWorkflowExecution\"} > 0.33",
        duracion="10m",
        caveats="Percentiles precalculados 1 min: no se reagregan a ventanas largas. Ns de poco tráfico: p95 refleja un solo request lento. Gate obligatorio.",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="No (añadir con gate)",
        dashboard="D01",
        pareja="temporal_cloud_v1_service_request_count",
    ),
    dict(
        metrica="temporal_cloud_v1_service_latency_p99",
        categoria="Frontend Service",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 99th percentile latency of service requests in seconds.",
        desc_ops="Más ruidoso que p95. Preferir p95 para alerta. p99 en tablero.",
        labels="operation",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta salvo SLO p99 explícito y alto volumen.",
        promql="—",
        duracion="—",
        caveats="Mismas reglas de no-agregación y samples mínimos que p95.",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="No",
        dashboard="D01",
        pareja="temporal_cloud_v1_service_request_count",
    ),
    # --- Workflow completion ---
    dict(
        metrica="temporal_cloud_v1_workflow_success_count",
        categoria="Workflow Completion",
        tipo="Rate",
        unidad="completions/s",
        desc_oficial="Successful workflow completions per second.",
        desc_ops="Denominador del fail rate. Tablero de throughput. No alertar éxitos a la baja sin contexto (puede ser menos negocio).",
        labels="temporal_task_queue, temporal_workflow_type",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Usar como denominador de fail+timeout.",
        promql="—",
        duracion="—",
        caveats="Alta cardinalidad por workflow type × task queue.",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="Sí (denominador FailRate)",
        dashboard="D01, D02",
        pareja="temporal_cloud_v1_workflow_failed_count, temporal_cloud_v1_workflow_timeout_count",
    ),
    dict(
        metrica="temporal_cloud_v1_workflow_failed_count",
        categoria="Workflow Completion",
        tipo="Rate",
        unidad="failures/s",
        desc_oficial="Workflow failures per second.",
        desc_ops="Fallo de negocio/aplicación visible en Cloud. Alertar ratio fail+timeout / (success+fail+timeout) por temporal_workflow_type crítico, no el absoluto global.",
        labels="temporal_task_queue, temporal_workflow_type",
        optin="—",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="(fail+timeout)/(success+fail+timeout) > 5% ≥15m, restringido a workflow types de producción.",
        promql="(sum by (temporal_namespace, temporal_workflow_type) (temporal_cloud_v1_workflow_failed_count) + sum by (temporal_namespace, temporal_workflow_type) (temporal_cloud_v1_workflow_timeout_count)) / clamp_min(sum by (temporal_namespace, temporal_workflow_type) (temporal_cloud_v1_workflow_success_count + temporal_cloud_v1_workflow_failed_count + temporal_cloud_v1_workflow_timeout_count), 0.001) > 0.05",
        duracion="15m",
        caveats="Alta cardinalidad. No paging global: un type de prueba ruidoso ensucia el ns. No es el worker caído (eso es no_poller).",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="Sí (TemporalCloudWorkflowFailRate)",
        dashboard="D01, D02",
        pareja="temporal_cloud_v1_workflow_success_count, temporal_cloud_v1_workflow_timeout_count",
    ),
    dict(
        metrica="temporal_cloud_v1_workflow_timeout_count",
        categoria="Workflow Completion",
        tipo="Rate",
        unidad="timeouts/s",
        desc_oficial="Workflow timeouts per second.",
        desc_ops="Workflows que llegaron a ExecutionTimeout/RunTimeout. Suele ser worker lento, activity colgada o timeout mal dimensionado. Sumar al fail rate.",
        labels="temporal_task_queue, temporal_workflow_type",
        optin="—",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="Parte del ratio fail+timeout. Umbral absoluto > 0 solo en types críticos.",
        promql="ver workflow_failed_count (ratio combinado)",
        duracion="15m",
        caveats="Timeout de workflow ≠ timeout de activity. Cruzar con activity_timeout_count y backlog_age.",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="Sí (TemporalCloudWorkflowFailRate)",
        dashboard="D02",
        pareja="temporal_cloud_v1_workflow_failed_count",
    ),
    dict(
        metrica="temporal_cloud_v1_workflow_cancel_count",
        categoria="Workflow Completion",
        tipo="Rate",
        unidad="cancellations/s",
        desc_oficial="Workflow cancellations per second.",
        desc_ops="A menudo operador o overlap policy. Tablero; alerta solo si cancel no es un camino de negocio.",
        labels="temporal_task_queue, temporal_workflow_type",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta por defecto.",
        promql="—",
        duracion="—",
        caveats="Cancel no es fail. No mezclar en el fail rate.",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="No",
        dashboard="D01",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_workflow_terminate_count",
        categoria="Workflow Completion",
        tipo="Rate",
        unidad="terminations/s",
        desc_oficial="Workflow terminations per second.",
        desc_ops="Terminate es acción de operador/admin. Pico = runbook o script. No es SLO de aplicación.",
        labels="temporal_task_queue, temporal_workflow_type",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Opcional: > 0 inesperado en prod (cambio no autorizado).",
        promql="—",
        duracion="—",
        caveats="No es fallo de Temporal Cloud.",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="No",
        dashboard="D01",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_workflow_continued_as_new_count",
        categoria="Workflow Completion",
        tipo="Rate",
        unidad="continue-as-new/s",
        desc_oficial="Workflows continued as new per second.",
        desc_ops="Esperado en workflows largos. Subida súbita puede ser history size / Continue-As-New mal puesto. Tablero.",
        labels="temporal_task_queue, temporal_workflow_type",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta por defecto.",
        promql="—",
        duracion="—",
        caveats="No es error. Cada Continue-As-New cierra un run y abre otro (Actions extra).",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="No",
        dashboard="D01",
        pareja="temporal_cloud_v1_billable_action_count",
    ),
    dict(
        metrica="temporal_cloud_v1_workflow_schedule_to_close_latency_p50",
        categoria="Workflow Completion",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 50th percentile workflow schedule-to-close latency in seconds.",
        desc_ops="Duración E2E del workflow (mediana). Tablero de negocio, no alerta de plataforma.",
        labels="temporal_workflow_type",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="SLO de negocio por type, no paging.",
        promql="—",
        duracion="—",
        caveats="No agregar entre workflow types. No incluye workflows aún abiertos.",
        cardinalidad="Media",
        requiere="—",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_workflow_schedule_to_close_latency_p95",
        categoria="Workflow Completion",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 95th percentile workflow schedule-to-close latency in seconds.",
        desc_ops="SLO de duración por workflow type crítico. Alerta de negocio (P2) si hay umbral de producto, con gate de completions.",
        labels="temporal_workflow_type",
        optin="—",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="Solo types con SLA. p95 > SLA AND (success+fail+timeout) con samples suficientes. No avg entre types.",
        promql="temporal_cloud_v1_workflow_schedule_to_close_latency_p95{temporal_workflow_type=\"<CRITICO>\"} > <SLA_s>",
        duracion="15m",
        caveats="Percentil no agregable. Workflows abiertos no entran. Umbral es de producto, no de Cloud.",
        cardinalidad="Media",
        requiere="SLA de negocio",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_workflow_success_count",
    ),
    dict(
        metrica="temporal_cloud_v1_workflow_schedule_to_close_latency_p99",
        categoria="Workflow Completion",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 99th percentile workflow schedule-to-close latency in seconds.",
        desc_ops="Cola del E2E. Tablero. Preferir p95 para alerta.",
        labels="temporal_workflow_type",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta por defecto.",
        promql="—",
        duracion="—",
        caveats="No agregar entre dimensiones.",
        cardinalidad="Media",
        requiere="—",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    # --- Activity ---
    dict(
        metrica="temporal_cloud_v1_activity_success_count",
        categoria="Activity",
        tipo="Rate",
        unidad="completions/s",
        desc_oficial="Successful activity completions per second.",
        desc_ops="Denominador del fail rate de activities. Standalone Activities usan temporal_workflow_type=__standalone_activity.",
        labels="temporal_task_queue, temporal_workflow_type",
        optin="temporal_activity_type",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Denominador. No habilitar temporal_activity_type en scrape 24×7 salvo necesidad.",
        promql="—",
        duracion="—",
        caveats="Alta cardinalidad si se activa activity_type. Label opt-in no viene por defecto.",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_activity_fail_count",
    ),
    dict(
        metrica="temporal_cloud_v1_activity_fail_count",
        categoria="Activity",
        tipo="Rate",
        unidad="failures/s",
        desc_oficial="Activity failures per second.",
        desc_ops="Fallo de activity (aplicación/dependencia). Candidata P2 por workflow type crítico. Causa típica: core, timeout corto, Key Vault. No es Cloud caído.",
        labels="temporal_task_queue, temporal_workflow_type",
        optin="temporal_activity_type",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="fail / (success+fail) > umbral en types críticos. Preferir SDK metrics para la causa.",
        promql="sum by (temporal_namespace, temporal_workflow_type) (temporal_cloud_v1_activity_fail_count) / clamp_min(sum by (temporal_namespace, temporal_workflow_type) (temporal_cloud_v1_activity_success_count + temporal_cloud_v1_activity_fail_count), 0.001) > 0.05",
        duracion="15m",
        caveats="Cloud no distingue 'error de negocio' vs 'timeout de dependency'. temporal_activity_type es opt-in.",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="No (candidato nuevo)",
        dashboard="—",
        pareja="temporal_cloud_v1_activity_success_count, temporal_cloud_v1_activity_timeout_count",
    ),
    dict(
        metrica="temporal_cloud_v1_activity_timeout_count",
        categoria="Activity",
        tipo="Rate",
        unidad="timeouts/s",
        desc_oficial="Activity timeouts per second.",
        desc_ops="StartToClose / ScheduleToClose / ScheduleToStart / Heartbeat. Label timeout_type. P2: activities que no terminan a tiempo (worker saturado o dependency lenta).",
        labels="temporal_task_queue, temporal_workflow_type, timeout_type",
        optin="temporal_activity_type",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="> 0 sostenido en types de prod, o ratio vs success. Filtrar timeout_type.",
        promql="sum by (temporal_namespace, timeout_type) (temporal_cloud_v1_activity_timeout_count) > 0",
        duracion="10m",
        caveats="ScheduleToStart alto = matching/backlog (cruzar no_poller y backlog_age). StartToClose alto = ejecución lenta.",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="No (candidato nuevo)",
        dashboard="—",
        pareja="temporal_cloud_v1_approximate_backlog_age",
    ),
    dict(
        metrica="temporal_cloud_v1_activity_task_fail_count",
        categoria="Activity",
        tipo="Rate",
        unidad="task failures/s",
        desc_oficial="Activity task failures per second.",
        desc_ops="Fallo al procesar el task (worker/SDK), distinto de activity_fail (resultado de negocio). P2 si sube: panic, serialization, worker crash mid-task.",
        labels="temporal_task_queue, temporal_workflow_type",
        optin="temporal_activity_type",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="> 0 sostenido en prod. Cruzar con SDK worker_task_failure.",
        promql="sum by (temporal_namespace, temporal_task_queue) (temporal_cloud_v1_activity_task_fail_count) > 0",
        duracion="10m",
        caveats="No confundir con activity_fail_count. Task fail suele reentregar el task.",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="No (candidato nuevo)",
        dashboard="—",
        pareja="temporal_cloud_v1_activity_fail_count",
    ),
    dict(
        metrica="temporal_cloud_v1_activity_task_timeout_count",
        categoria="Activity",
        tipo="Rate",
        unidad="task timeouts/s",
        desc_oficial="Activity task timeouts per second.",
        desc_ops="El worker no completó el task a tiempo (heartbeat perdido, GC, kill). P2. timeout_type aplica.",
        labels="temporal_task_queue, temporal_workflow_type, timeout_type",
        optin="temporal_activity_type",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="> 0 sostenido. Suele ir con backlog_age y SDK schedule_to_start.",
        promql="sum by (temporal_namespace, temporal_task_queue) (temporal_cloud_v1_activity_task_timeout_count) > 0",
        duracion="10m",
        caveats="Distinto de activity_timeout_count (timeouts de la Activity Execution).",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="No (candidato nuevo)",
        dashboard="—",
        pareja="temporal_cloud_v1_activity_timeout_count",
    ),
    dict(
        metrica="temporal_cloud_v1_activity_cancel_count",
        categoria="Activity",
        tipo="Rate",
        unidad="cancellations/s",
        desc_oficial="Activity cancellations per second.",
        desc_ops="Cancel de activity (workflow canceló o timeout de padre). Tablero.",
        labels="temporal_task_queue, temporal_workflow_type",
        optin="temporal_activity_type",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta por defecto.",
        promql="—",
        duracion="—",
        caveats="Cancel no es fail.",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_activity_terminate_count",
        categoria="Activity",
        tipo="Rate",
        unidad="terminations/s",
        desc_oficial="Activity terminations per second. This metric only applies to Standalone Activities. Regular Activities that run within a Workflow cannot be terminated independently.",
        desc_ops="Solo Standalone Activities. Si no usan ActivityClient, la serie estará vacía.",
        labels="temporal_task_queue, temporal_workflow_type",
        optin="temporal_activity_type",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Solo si hay Standalone Activities en prod y terminate no es esperado.",
        promql="—",
        duracion="—",
        caveats="temporal_workflow_type = __standalone_activity. No aplica a activities dentro de workflow.",
        cardinalidad="Alta",
        requiere="Standalone Activities",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_activity_start_to_close_latency_p50",
        categoria="Activity",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 50th percentile activity start-to-close latency in seconds.",
        desc_ops="Tiempo de ejecución (sin cola). Tablero. Labels de latencia: solo temporal_activity_type (opt-in).",
        labels="— (latencia: solo activity_type)",
        optin="temporal_activity_type",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta p50.",
        promql="—",
        duracion="—",
        caveats="Sin task_queue ni workflow_type a propósito: percentiles no se agregan. Habilitar activity_type solo si hace falta.",
        cardinalidad="Media (si opt-in)",
        requiere="—",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_activity_start_to_close_latency_p95",
        categoria="Activity",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 95th percentile activity start-to-close latency in seconds.",
        desc_ops="Candidata P2 si hay SLA de ejecución de activity y temporal_activity_type está en el scrape. Si no, usar SDK.",
        labels="— (latencia: solo activity_type)",
        optin="temporal_activity_type",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="Solo con label opt-in y SLA. Sin activity_type la serie es un solo p95 mezclado: no alertar.",
        promql="temporal_cloud_v1_activity_start_to_close_latency_p95{temporal_activity_type=\"<CRITICA>\"} > <SLA_s>",
        duracion="15m",
        caveats="Requiere ?labels=temporal_activity_type. No mezclar con schedule-to-close (incluye cola).",
        cardinalidad="Media (si opt-in)",
        requiere="label opt-in + SLA",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_activity_schedule_to_close_latency_p95",
    ),
    dict(
        metrica="temporal_cloud_v1_activity_start_to_close_latency_p99",
        categoria="Activity",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 99th percentile activity start-to-close latency in seconds.",
        desc_ops="Cola de ejecución. Tablero. Preferir p95.",
        labels="— (latencia: solo activity_type)",
        optin="temporal_activity_type",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta por defecto.",
        promql="—",
        duracion="—",
        caveats="No agregar entre dimensiones.",
        cardinalidad="Media (si opt-in)",
        requiere="—",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_activity_schedule_to_close_latency_p50",
        categoria="Activity",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 50th percentile activity schedule-to-close latency in seconds.",
        desc_ops="Cola + ejecución (mediana). Tablero.",
        labels="— (latencia: solo activity_type)",
        optin="temporal_activity_type",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta p50.",
        promql="—",
        duracion="—",
        caveats="Incluye tiempo en task queue. Separar de start-to-close para ver matching vs worker.",
        cardinalidad="Media (si opt-in)",
        requiere="—",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_approximate_backlog_age",
    ),
    dict(
        metrica="temporal_cloud_v1_activity_schedule_to_close_latency_p95",
        categoria="Activity",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 95th percentile activity schedule-to-close latency in seconds.",
        desc_ops="SLA E2E de activity (espera + run). P2 si hay umbral y label opt-in. Si p95 schedule>>start, el problema es cola (backlog/no_poller).",
        labels="— (latencia: solo activity_type)",
        optin="temporal_activity_type",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="Con activity_type y SLA. Comparar vs start-to-close p95.",
        promql="temporal_cloud_v1_activity_schedule_to_close_latency_p95{temporal_activity_type=\"<CRITICA>\"} > <SLA_s>",
        duracion="15m",
        caveats="Sin opt-in no desglosa. Preferible cruzar backlog_age (no necesita opt-in).",
        cardinalidad="Media (si opt-in)",
        requiere="label opt-in + SLA",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_activity_start_to_close_latency_p95",
    ),
    dict(
        metrica="temporal_cloud_v1_activity_schedule_to_close_latency_p99",
        categoria="Activity",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 99th percentile activity schedule-to-close latency in seconds.",
        desc_ops="Tablero. Preferir p95.",
        labels="— (latencia: solo activity_type)",
        optin="temporal_activity_type",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta por defecto.",
        promql="—",
        duracion="—",
        caveats="No agregar entre dimensiones.",
        cardinalidad="Media (si opt-in)",
        requiere="—",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    # --- Nexus (Public Preview) ---
    dict(
        metrica="temporal_cloud_v1_nexus_operation_schedule_to_close_latency_p50",
        categoria="Nexus (Public Preview)",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 50th percentile time from when a Nexus Operation is scheduled until it reaches a terminal state, in seconds. Includes sync and async completions.",
        desc_ops="Caller side. temporal_namespace = caller. Tablero si hay Nexus.",
        labels="nexus_endpoint, temporal_workflow_type",
        optin="nexus_service, nexus_operation",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Ignorar si el ns no usa Nexus.",
        promql="—",
        duracion="—",
        caveats="Public Preview. Alta cardinalidad. No filtrar sync vs async. Standalone: workflow_type=__temporal_standalone_nexus_operation__.",
        cardinalidad="Alta",
        requiere="Nexus",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_operation_schedule_to_close_latency_p95",
        categoria="Nexus (Public Preview)",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 95th percentile schedule-to-close latency of a Nexus Operation, in seconds.",
        desc_ops="P2 solo si Nexus está en el camino crítico. No agregar percentiles.",
        labels="nexus_endpoint, temporal_workflow_type",
        optin="nexus_service, nexus_operation",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="Solo namespaces con Nexus y SLA. Si no hay Nexus, no crear la alerta.",
        promql="temporal_cloud_v1_nexus_operation_schedule_to_close_latency_p95 > <SLA_s>",
        duracion="15m",
        caveats="Public Preview. Mezcla sync y async.",
        cardinalidad="Alta",
        requiere="Nexus + SLA",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_nexus_operation_fail_count",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_operation_schedule_to_close_latency_p99",
        categoria="Nexus (Public Preview)",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 99th percentile schedule-to-close latency of a Nexus Operation, in seconds.",
        desc_ops="Tablero Nexus. Preferir p95.",
        labels="nexus_endpoint, temporal_workflow_type",
        optin="nexus_service, nexus_operation",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Ignorar sin Nexus.",
        promql="—",
        duracion="—",
        caveats="Public Preview.",
        cardinalidad="Alta",
        requiere="Nexus",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_operation_schedule_to_start_latency_p50",
        categoria="Nexus (Public Preview)",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 50th percentile time from schedule until async start. For sync/close-without-start equals schedule-to-close.",
        desc_ops="Espera hasta start async. Tablero.",
        labels="nexus_endpoint, temporal_workflow_type",
        optin="nexus_service, nexus_operation",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Ignorar sin Nexus.",
        promql="—",
        duracion="—",
        caveats="No distingue sync vs async en las labels.",
        cardinalidad="Alta",
        requiere="Nexus",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_operation_schedule_to_start_latency_p95",
        categoria="Nexus (Public Preview)",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 95th percentile schedule-to-start latency of a Nexus Operation, in seconds.",
        desc_ops="P2 si el start async es el SLO (handler no arranca). Solo con Nexus.",
        labels="nexus_endpoint, temporal_workflow_type",
        optin="nexus_service, nexus_operation",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="Namespaces Nexus + SLA de arranque.",
        promql="temporal_cloud_v1_nexus_operation_schedule_to_start_latency_p95 > <SLA_s>",
        duracion="15m",
        caveats="Public Preview. Sync mezcla el valor con schedule-to-close.",
        cardinalidad="Alta",
        requiere="Nexus + SLA",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_operation_schedule_to_start_latency_p99",
        categoria="Nexus (Public Preview)",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 99th percentile schedule-to-start latency of a Nexus Operation, in seconds.",
        desc_ops="Tablero. Preferir p95.",
        labels="nexus_endpoint, temporal_workflow_type",
        optin="nexus_service, nexus_operation",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Ignorar sin Nexus.",
        promql="—",
        duracion="—",
        caveats="Public Preview.",
        cardinalidad="Alta",
        requiere="Nexus",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_operation_start_to_close_latency_p50",
        categoria="Nexus (Public Preview)",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 50th percentile time from async start until terminal state. Not recorded for sync or close-without-start.",
        desc_ops="Duración async. Tablero.",
        labels="nexus_endpoint, temporal_workflow_type",
        optin="nexus_service, nexus_operation",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Ignorar sin Nexus.",
        promql="—",
        duracion="—",
        caveats="Vacío si todas las ops son síncronas.",
        cardinalidad="Alta",
        requiere="Nexus async",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_operation_start_to_close_latency_p95",
        categoria="Nexus (Public Preview)",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 95th percentile start-to-close latency of an async Nexus Operation, in seconds.",
        desc_ops="P2 si el handler async tiene SLA. No se emite en sync.",
        labels="nexus_endpoint, temporal_workflow_type",
        optin="nexus_service, nexus_operation",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="Nexus async + SLA.",
        promql="temporal_cloud_v1_nexus_operation_start_to_close_latency_p95 > <SLA_s>",
        duracion="15m",
        caveats="Public Preview. Ausente en operaciones síncronas.",
        cardinalidad="Alta",
        requiere="Nexus async + SLA",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_operation_start_to_close_latency_p99",
        categoria="Nexus (Public Preview)",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 99th percentile start-to-close latency of an async Nexus Operation, in seconds.",
        desc_ops="Tablero. Preferir p95.",
        labels="nexus_endpoint, temporal_workflow_type",
        optin="nexus_service, nexus_operation",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Ignorar sin Nexus.",
        promql="—",
        duracion="—",
        caveats="Public Preview.",
        cardinalidad="Alta",
        requiere="Nexus async",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_operation_success_count",
        categoria="Nexus (Public Preview)",
        tipo="Rate",
        unidad="ops/s",
        desc_oficial="Successfully completed Nexus Operations per second.",
        desc_ops="Denominador caller-side. temporal_namespace = caller.",
        labels="nexus_endpoint, temporal_workflow_type",
        optin="nexus_service, nexus_operation",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Denominador del fail rate Nexus.",
        promql="—",
        duracion="—",
        caveats="Public Preview.",
        cardinalidad="Alta",
        requiere="Nexus",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_nexus_operation_fail_count",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_operation_cancel_count",
        categoria="Nexus (Public Preview)",
        tipo="Rate",
        unidad="ops/s",
        desc_oficial="Canceled Nexus Operations per second.",
        desc_ops="Tablero. No mezclar con fail.",
        labels="nexus_endpoint, temporal_workflow_type",
        optin="nexus_service, nexus_operation",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Ignorar sin Nexus.",
        promql="—",
        duracion="—",
        caveats="Public Preview.",
        cardinalidad="Alta",
        requiere="Nexus",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_operation_terminate_count",
        categoria="Nexus (Public Preview)",
        tipo="Rate",
        unidad="ops/s",
        desc_oficial="Nexus Operations terminated before successful completion per second.",
        desc_ops="Tablero / forense. No paging por defecto.",
        labels="nexus_endpoint, temporal_workflow_type",
        optin="nexus_service, nexus_operation",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Ignorar sin Nexus.",
        promql="—",
        duracion="—",
        caveats="Public Preview.",
        cardinalidad="Alta",
        requiere="Nexus",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_operation_fail_count",
        categoria="Nexus (Public Preview)",
        tipo="Rate",
        unidad="ops/s",
        desc_oficial="Failed Nexus Operations per second.",
        desc_ops="P2 si Nexus está en producción. Ratio vs success. Caller namespace.",
        labels="nexus_endpoint, temporal_workflow_type",
        optin="nexus_service, nexus_operation",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="Solo si hay Nexus. fail/(success+fail) > umbral.",
        promql="sum by (temporal_namespace, nexus_endpoint) (temporal_cloud_v1_nexus_operation_fail_count) > 0",
        duracion="10m",
        caveats="Public Preview. No crear la regla 'por si acaso' en ns sin Nexus.",
        cardinalidad="Alta",
        requiere="Nexus",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_nexus_operation_success_count",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_operation_timeout_count",
        categoria="Nexus (Public Preview)",
        tipo="Rate",
        unidad="ops/s",
        desc_oficial="Nexus Operations that timed out before completing execution per second.",
        desc_ops="P2 con Nexus. Label timeout_type.",
        labels="nexus_endpoint, temporal_workflow_type, timeout_type",
        optin="nexus_service, nexus_operation",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="> 0 sostenido en ns con Nexus.",
        promql="sum by (temporal_namespace, nexus_endpoint) (temporal_cloud_v1_nexus_operation_timeout_count) > 0",
        duracion="10m",
        caveats="Public Preview.",
        cardinalidad="Alta",
        requiere="Nexus",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_nexus_operation_fail_count",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_request_count",
        categoria="Nexus (Public Preview)",
        tipo="Rate",
        unidad="requests/s",
        desc_oficial="Nexus requests received by the Nexus Endpoint per second.",
        desc_ops="Handler side. temporal_namespace = handler. Denominador de request_error.",
        labels="nexus_endpoint, nexus_caller_type",
        optin="nexus_service, nexus_operation",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Denominador. Distinto de operation_* (caller).",
        promql="—",
        duracion="—",
        caveats="Public Preview. No confundir caller vs handler namespace.",
        cardinalidad="Alta",
        requiere="Nexus handler",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_nexus_request_error_count",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_request_error_count",
        categoria="Nexus (Public Preview)",
        tipo="Rate",
        unidad="errors/s",
        desc_oficial="Failed Nexus requests per second.",
        desc_ops="P2 handler-side. El endpoint no procesa el request. Distinto de operation_fail (resultado de la operación).",
        labels="nexus_endpoint, nexus_caller_type",
        optin="nexus_service, nexus_operation",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="error/request > umbral en ns handler.",
        promql="sum by (temporal_namespace, nexus_endpoint) (temporal_cloud_v1_nexus_request_error_count) / clamp_min(sum by (temporal_namespace, nexus_endpoint) (temporal_cloud_v1_nexus_request_count), 0.001) > 0.01",
        duracion="10m",
        caveats="Public Preview. Namespace = handler, no caller.",
        cardinalidad="Alta",
        requiere="Nexus handler",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_nexus_request_count",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_request_latency_p50",
        categoria="Nexus (Public Preview)",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 50th percentile time the handler takes to process one Nexus request. For async start, ends when the handler finishes the start request, not when the Operation completes.",
        desc_ops="Latencia del handler por request (no E2E de la operación). Tablero.",
        labels="nexus_endpoint, nexus_caller_type",
        optin="nexus_service, nexus_operation",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Ignorar sin Nexus.",
        promql="—",
        duracion="—",
        caveats="No es schedule-to-close de la operación.",
        cardinalidad="Alta",
        requiere="Nexus handler",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_request_latency_p95",
        categoria="Nexus (Public Preview)",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 95th percentile handler request latency, in seconds.",
        desc_ops="P2 de handler si hay SLA de request (no de operación completa).",
        labels="nexus_endpoint, nexus_caller_type",
        optin="nexus_service, nexus_operation",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="Nexus handler + SLA de request. Gate de nexus_request_count.",
        promql="temporal_cloud_v1_nexus_request_latency_p95 > <SLA_s>",
        duracion="10m",
        caveats="Public Preview. Async start no espera el completion.",
        cardinalidad="Alta",
        requiere="Nexus handler + SLA",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_nexus_request_count",
    ),
    dict(
        metrica="temporal_cloud_v1_nexus_request_latency_p99",
        categoria="Nexus (Public Preview)",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 99th percentile handler request latency, in seconds.",
        desc_ops="Tablero. Preferir p95.",
        labels="nexus_endpoint, nexus_caller_type",
        optin="nexus_service, nexus_operation",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Ignorar sin Nexus.",
        promql="—",
        duracion="—",
        caveats="Public Preview.",
        cardinalidad="Alta",
        requiere="Nexus handler",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    # --- Task Queue ---
    dict(
        metrica="temporal_cloud_v1_approximate_backlog_count",
        categoria="Task Queue",
        tipo="Value",
        unidad="tasks",
        desc_oficial="The approximate number of tasks pending in a task queue. Started Activities are not included (already dequeued).",
        desc_ops="Cola de matching. P2 con umbral por TQ. Aproximada: puede sobrecontar tasks inválidas y resetear a 0 si el TQ se descarga (~5 min idle).",
        labels="temporal_task_queue, task_type (workflow|activity), task_priority",
        optin="temporal_worker_deployment_name, temporal_worker_build_id",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="> umbral (p. ej. 100) ≥15m por ns + task_queue + task_type. Cruzar con backlog_age y no_poller.",
        promql="sum by (temporal_namespace, temporal_task_queue, task_type) (temporal_cloud_v1_approximate_backlog_count) > 100",
        duracion="15m",
        caveats="Oficial: overcount de tasks expiradas/canceladas hasta la cabeza de la cola. Reset a 0 en TQ idle ~5 min. No es exacto. Started activities no cuentan.",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="Sí (TemporalCloudBacklogHigh)",
        dashboard="D01, D02",
        pareja="temporal_cloud_v1_approximate_backlog_age, temporal_cloud_v1_no_poller_tasks_count",
    ),
    dict(
        metrica="temporal_cloud_v1_approximate_backlog_age",
        categoria="Task Queue",
        tipo="Value",
        unidad="seconds",
        desc_oficial="The approximate age, in seconds, of the oldest task in a task queue backlog.",
        desc_ops="Mejor señal de 'algo espera' que el count. P2 con umbral en segundos (p. ej. 60s). Mismas limitaciones de accuracy que count.",
        labels="temporal_task_queue, task_type (workflow|activity)",
        optin="temporal_worker_deployment_name, temporal_worker_build_id",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="age > 60s (ajustar) ≥10m. Preferible al count cuando el umbral de tasks es difícil.",
        promql="max by (temporal_namespace, temporal_task_queue, task_type) (temporal_cloud_v1_approximate_backlog_age) > 60",
        duracion="10m",
        caveats="Aproximada. Mismo reset a 0 en TQ idle. No está aún en alertas-catalogo.yaml → candidato nuevo.",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="No (candidato nuevo; mejor que count)",
        dashboard="—",
        pareja="temporal_cloud_v1_approximate_backlog_count",
    ),
    dict(
        metrica="temporal_cloud_v1_poll_success_count",
        categoria="Task Queue",
        tipo="Rate",
        unidad="matches/s",
        desc_oficial="Successfully matched tasks per second.",
        desc_ops="Matching sano. Denominador de sync match ratio. No alertar a la baja (menos negocio).",
        labels="operation, task_type, temporal_task_queue",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Tablero: poll_success_sync / poll_success = fracción sin espera.",
        promql="—",
        duracion="—",
        caveats="Alta cardinalidad por TQ.",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="No",
        dashboard="D01",
        pareja="temporal_cloud_v1_poll_success_sync_count",
    ),
    dict(
        metrica="temporal_cloud_v1_poll_success_sync_count",
        categoria="Task Queue",
        tipo="Rate",
        unidad="matches/s",
        desc_oficial="Tasks matched synchronously per second (no polling wait).",
        desc_ops="Match inmediato (task ya esperaba poller, o poller ya esperaba task). Ratio alto = matching caliente. Tablero.",
        labels="operation, task_type, temporal_task_queue",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta. Caída del ratio sugiere cola o pollers fríos.",
        promql="—",
        duracion="—",
        caveats="—",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="No",
        dashboard="D01",
        pareja="temporal_cloud_v1_poll_success_count",
    ),
    dict(
        metrica="temporal_cloud_v1_poll_timeout_count",
        categoria="Task Queue",
        tipo="Rate",
        unidad="poll timeouts/s",
        desc_oficial="The rate of poll requests that timed out without receiving a task.",
        desc_ops="Normal cuando hay pollers idle (long-poll vacío). NO alertar > 0. Un ratio timeout/(success+timeout) alto = workers de más, no incidente.",
        labels="operation, task_type, temporal_task_queue",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta. El overview usa success/(success+timeout) como occupancy, no como SLO.",
        promql="—",
        duracion="—",
        caveats="Timeout de poll es el caso feliz de un worker sin trabajo. Alertarlo genera ruido masivo.",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="No",
        dashboard="D01",
        pareja="temporal_cloud_v1_poll_success_count",
    ),
    dict(
        metrica="temporal_cloud_v1_no_poller_tasks_count",
        categoria="Task Queue",
        tipo="Rate",
        unidad="tasks/s",
        desc_oficial="The rate of tasks added to queues with no active pollers.",
        desc_ops="La señal más clara de worker ausente/mal task queue. P1 en prod. Cloud matching; la causa está en el worker (deploy, TQ mal escrito, pollers=0).",
        labels="temporal_task_queue, task_type",
        optin="—",
        importancia="P1",
        orden=1,
        candidata="Sí",
        severidad="P1",
        condicion="> 0 ≥3m en namespaces de producción.",
        promql="sum by (temporal_namespace, temporal_task_queue, task_type) (temporal_cloud_v1_no_poller_tasks_count) > 0",
        duracion="3m",
        caveats="No es un rebalance de brokers. Profundizar con SDK (pollers, last poll). TQ de test puede ruidar: filtrar ns prod.",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="Sí (TemporalCloudNoPollerTasks)",
        dashboard="D02",
        pareja="temporal_cloud_v1_approximate_backlog_count",
    ),
    # --- Namespace ---
    dict(
        metrica="temporal_cloud_v1_namespace_open_workflows",
        categoria="Namespace",
        tipo="Value",
        unidad="workflows",
        desc_oficial="The current number of open workflows in a namespace.",
        desc_ops="Inventario. P2 si salta vs 24h (fuga: workflows que no cierran). No alertar el absoluto sin baseline.",
        labels="— (solo comunes)",
        optin="—",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="valor > 1.5 × offset 1d, ≥30m.",
        promql="sum by (temporal_namespace) (temporal_cloud_v1_namespace_open_workflows) > 1.5 * sum by (temporal_namespace) (temporal_cloud_v1_namespace_open_workflows offset 1d)",
        duracion="30m",
        caveats="Gauge. Un batch legitimo también sube. Cruzar con success/fail y schedules.",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="Sí (TemporalCloudOpenWorkflowsJump)",
        dashboard="D01, D02",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_total_action_count",
        categoria="Namespace",
        tipo="Rate",
        unidad="actions/s",
        desc_oficial="The total number of actions performed per second. Actions with is_background=false count toward action_limit.",
        desc_ops="Cuota operativa APS, no factura. P3: is_background=false / action_limit > 80%. Distinto de billable_action_count.",
        labels="is_background, namespace_mode (active|standby)",
        optin="—",
        importancia="P3",
        orden=3,
        candidata="Sí",
        severidad="P3",
        condicion="total_action{is_background=false} / action_limit > 0.8 ≥15m.",
        promql='sum by (temporal_namespace) (temporal_cloud_v1_total_action_count{is_background="false"}) / sum by (temporal_namespace) (temporal_cloud_v1_action_limit) > 0.8',
        duracion="15m",
        caveats="NO incluye label region (scope namespace). Background no topa el limit. El % medio puede estar verde con throttle por burst (1 min).",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="Sí (TemporalCloudActionLimitHigh)",
        dashboard="D01, D03",
        pareja="temporal_cloud_v1_action_limit, temporal_cloud_v1_total_action_throttled_count",
    ),
    dict(
        metrica="temporal_cloud_v1_billable_action_count",
        categoria="Namespace",
        tipo="Rate",
        unidad="billable actions/s",
        desc_oficial="The number of billable actions per second, broken down by action_type and Workflow Type. Not all billable actions are included; estimate can differ by up to 5% from metering.",
        desc_ops="Señal FinOps. P3 vs baseline 7d. No paging. Top action_type: heartbeats y retry_activity suelen ser el ruido caro.",
        labels="action_type, temporal_workflow_type",
        optin="—",
        importancia="P3",
        orden=3,
        candidata="Sí",
        severidad="P3",
        condicion="rate actual > 1.5 × avg_over_time 7d AND > 0.1/s, ≥2h.",
        promql="sum by (temporal_namespace) (temporal_cloud_v1_billable_action_count) > 1.5 * sum by (temporal_namespace) (avg_over_time(temporal_cloud_v1_billable_action_count[7d]))",
        duracion="2h",
        caveats="Alta cardinalidad. Estimación ±5% vs factura. Excepciones oficiales (capacidad/export). Totales 7d/30d = sum_over_time * 60, no el invoice.",
        cardinalidad="Alta",
        requiere="—",
        en_yaml="Sí (TemporalCloudBillableVsBaseline7d)",
        dashboard="D03",
        pareja="temporal_cloud_v1_total_action_count",
    ),
    dict(
        metrica="temporal_cloud_v1_total_action_throttled_count",
        categoria="Namespace",
        tipo="Rate",
        unidad="throttled actions/s",
        desc_oficial="The total number of actions throttled per second. See Monitoring Trends Against Limits.",
        desc_ops="Cuota APS frenando trabajo. P1: > 0. El % vs limit puede estar < 80%.",
        labels="— (solo comunes; sin region en actions de ns — verificar scrape)",
        optin="—",
        importancia="P1",
        orden=1,
        candidata="Sí",
        severidad="P1",
        condicion="> 0 ≥2m.",
        promql="sum by (temporal_namespace) (temporal_cloud_v1_total_action_throttled_count) > 0",
        duracion="2m",
        caveats="Burst vs media 1 min. No sustituye resource_exhausted (otro mecanismo).",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="Sí (TemporalCloudThrottled)",
        dashboard="D02, D03",
        pareja="temporal_cloud_v1_action_limit, temporal_cloud_v1_total_action_count",
    ),
    dict(
        metrica="temporal_cloud_v1_operations_count",
        categoria="Namespace",
        tipo="Rate",
        unidad="operations/s",
        desc_oficial="Operations performed per second.",
        desc_ops="Otra cuota (operations, no Actions). P3 vs operations_limit. is_background=false cuenta al tope.",
        labels="operation, is_background, namespace_mode",
        optin="—",
        importancia="P3",
        orden=3,
        candidata="Sí",
        severidad="P3",
        condicion="operations{is_background=false} / operations_limit > 0.8.",
        promql='sum by (temporal_namespace) (temporal_cloud_v1_operations_count{is_background="false"}) / sum by (temporal_namespace) (temporal_cloud_v1_operations_limit) > 0.8',
        duracion="15m",
        caveats="No es billable. Namespace_mode standby en HA.",
        cardinalidad="Media",
        requiere="—",
        en_yaml="Sí (TemporalCloudOperationsLimitHigh)",
        dashboard="D01, D03",
        pareja="temporal_cloud_v1_operations_limit, temporal_cloud_v1_operations_throttled_count",
    ),
    dict(
        metrica="temporal_cloud_v1_operations_throttled_count",
        categoria="Namespace",
        tipo="Rate",
        unidad="throttled ops/s",
        desc_oficial="Operations throttled due to rate limits per second.",
        desc_ops="P1 junto a los otros throttle. Label namespace_mode indica active vs standby.",
        labels="operation, is_background, namespace_mode",
        optin="—",
        importancia="P1",
        orden=1,
        candidata="Sí",
        severidad="P1",
        condicion="> 0 ≥2m.",
        promql="sum by (temporal_namespace) (temporal_cloud_v1_operations_throttled_count) > 0",
        duracion="2m",
        caveats="Mismo caveat de burst 1 min.",
        cardinalidad="Media",
        requiere="—",
        en_yaml="Sí (TemporalCloudThrottled)",
        dashboard="D02",
        pareja="temporal_cloud_v1_operations_limit",
    ),
    # --- Schedule ---
    dict(
        metrica="temporal_cloud_v1_schedule_action_success_count",
        categoria="Schedule",
        tipo="Rate",
        unidad="executions/s",
        desc_oficial="Successful scheduled executions per second.",
        desc_ops="Denominador / tablero. Si el ns no usa Schedule, serie vacía (normal).",
        labels="— (solo comunes)",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta. Silencio de schedule se ve mejor con missed_catchup.",
        promql="—",
        duracion="—",
        caveats="No se emite si no hay Schedules.",
        cardinalidad="Baja",
        requiere="Schedule",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_schedule_missed_catchup_window_count",
    ),
    dict(
        metrica="temporal_cloud_v1_schedule_buffer_overruns_count",
        categoria="Schedule",
        tipo="Rate",
        unidad="dropped/s",
        desc_oficial="Scheduled executions dropped per second because the Schedule's internal Action buffer is full.",
        desc_ops="P2: se perdieron corridas. El schedule genera más rápido de lo que drena el buffer.",
        labels="— (solo comunes)",
        optin="—",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="> 0 ≥5m. Solo ns con Schedules.",
        promql="sum by (temporal_namespace) (temporal_cloud_v1_schedule_buffer_overruns_count) > 0",
        duracion="5m",
        caveats="Trabajo perdido (no retrasado). Revisar overlap policy y duración del workflow disparado.",
        cardinalidad="Baja",
        requiere="Schedule",
        en_yaml="Sí (TemporalCloudScheduleBufferOverrun)",
        dashboard="D02",
        pareja="temporal_cloud_v1_schedule_overlap_skipped_count",
    ),
    dict(
        metrica="temporal_cloud_v1_schedule_missed_catchup_window_count",
        categoria="Schedule",
        tipo="Rate",
        unidad="dropped/s",
        desc_oficial="Scheduled executions permanently dropped per second because they fell outside the Catchup Window. Not emitted for backfills.",
        desc_ops="P2: corridas fuera de catchup, no se recuperan. Outage de worker más largo que la ventana, o Cloud no pudo ponerse al día.",
        labels="— (solo comunes)",
        optin="—",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="> 0 ≥5m.",
        promql="sum by (temporal_namespace) (temporal_cloud_v1_schedule_missed_catchup_window_count) > 0",
        duracion="5m",
        caveats="No incluye backfills. Ajustar CatchupWindow del Schedule si el downtime esperado es mayor.",
        cardinalidad="Baja",
        requiere="Schedule",
        en_yaml="Sí (TemporalCloudScheduleMissedCatchup)",
        dashboard="D02",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_schedule_rate_limited_count",
        categoria="Schedule",
        tipo="Rate",
        unidad="delayed/s",
        desc_oficial="[DEPRECATED] Scheduled executions delayed per second due to rate limiting. Being replaced by standard request-per-second rate limit monitoring with dynamic scaling. Will be removed when rolled out to all of Temporal Cloud.",
        desc_ops="NO alerta nueva. Usar service_request_throttled / operations_throttled. Retirar si aún está en dashboards.",
        labels="— (solo comunes)",
        optin="—",
        importancia="No alerta",
        orden=5,
        candidata="No",
        severidad="—",
        condicion="Deprecada. No crear alerta.",
        promql="—",
        duracion="—",
        caveats="Oficialmente DEPRECATED. Desaparecerá del endpoint.",
        cardinalidad="Baja",
        requiere="Schedule (legado)",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_service_request_throttled_count",
    ),
    dict(
        metrica="temporal_cloud_v1_schedule_overlap_skipped_count",
        categoria="Schedule",
        tipo="Rate",
        unidad="skipped/s",
        desc_oficial="Schedule executions dropped per second due to the overlap policy. Emitted for SKIP when another execution is running/selected, and for BUFFER_ONE when its single buffered slot is occupied. Can also emit during backfills with SKIP/BUFFER_ONE. Sustained rate: executions generated faster than they complete.",
        desc_ops="P2: el schedule pisa sus propias corridas. Política SKIP/BUFFER_ONE. No es Cloud caído; es diseño o workflow demasiado lento.",
        labels="schedule_overlap_policy",
        optin="—",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="> 0 ≥10m.",
        promql="sum by (temporal_namespace) (temporal_cloud_v1_schedule_overlap_skipped_count) > 0",
        duracion="10m",
        caveats="Puede ser intencional (SKIP). Confirmar overlap policy antes de paging. Incluye backfills SKIP/BUFFER_ONE.",
        cardinalidad="Baja",
        requiere="Schedule",
        en_yaml="Sí (TemporalCloudScheduleOverlapSkipped)",
        dashboard="D02",
        pareja="temporal_cloud_v1_schedule_buffer_overruns_count",
    ),
    dict(
        metrica="temporal_cloud_v1_schedule_action_e2e_delay_p50",
        categoria="Schedule",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 50th percentile total delay from an execution's scheduled time to actual start, in seconds. Includes overlap-policy wait. Not emitted for backfills.",
        desc_ops="Retraso de arranque del Schedule. Tablero.",
        labels="— (solo comunes)",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta p50.",
        promql="—",
        duracion="—",
        caveats="No se emite en backfills. Incluye espera por overlap.",
        cardinalidad="Baja",
        requiere="Schedule",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_schedule_action_e2e_delay_p95",
        categoria="Schedule",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 95th percentile delay from scheduled time to actual start, in seconds.",
        desc_ops="P2 si el Schedule tiene SLA de puntualidad (p. ej. batch 02:00).",
        labels="— (solo comunes)",
        optin="—",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="p95 > tolerancia (p. ej. 60s) en ns con Schedules de negocio.",
        promql="temporal_cloud_v1_schedule_action_e2e_delay_p95 > 60",
        duracion="15m",
        caveats="No backfills. Gate: si no hay schedules, la serie falta (no alerta en ausencia).",
        cardinalidad="Baja",
        requiere="Schedule + SLA puntualidad",
        en_yaml="No (candidato nuevo)",
        dashboard="—",
        pareja="temporal_cloud_v1_schedule_overlap_skipped_count",
    ),
    dict(
        metrica="temporal_cloud_v1_schedule_action_e2e_delay_p99",
        categoria="Schedule",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 99th percentile delay from scheduled time to actual start, in seconds.",
        desc_ops="Tablero. Preferir p95.",
        labels="— (solo comunes)",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta por defecto.",
        promql="—",
        duracion="—",
        caveats="No backfills.",
        cardinalidad="Baja",
        requiere="Schedule",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    # --- Replication ---
    dict(
        metrica="temporal_cloud_v1_replication_lag_p50",
        categoria="Replication",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 50th percentile cross-region replication lag in seconds.",
        desc_ops="Solo namespaces HA. 0/ausente si no hay high availability. Tablero.",
        labels="— (solo comunes)",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta p50. Usar p99.",
        promql="—",
        duracion="—",
        caveats="Sin HA la serie no aplica. No crear la alerta 'por si acaso'.",
        cardinalidad="Baja",
        requiere="Namespace HA",
        en_yaml="No",
        dashboard="—",
        pareja="—",
    ),
    dict(
        metrica="temporal_cloud_v1_replication_lag_p95",
        categoria="Replication",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 95th percentile cross-region replication lag in seconds.",
        desc_ops="Tablero HA. Alertar p99 (cola de replicación).",
        labels="— (solo comunes)",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Preferir p99.",
        promql="—",
        duracion="—",
        caveats="Solo HA.",
        cardinalidad="Baja",
        requiere="Namespace HA",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_replication_lag_p99",
    ),
    dict(
        metrica="temporal_cloud_v1_replication_lag_p99",
        categoria="Replication",
        tipo="Latency",
        unidad="seconds",
        desc_oficial="The 99th percentile cross-region replication lag in seconds.",
        desc_ops="P2 en namespaces HA. Lag alto = failover con pérdida de recencia. Desactivar la regla si no hay HA.",
        labels="— (solo comunes)",
        optin="—",
        importancia="P2",
        orden=2,
        candidata="Sí",
        severidad="P2",
        condicion="p99 > 30s ≥10m, solo ns HA.",
        promql="temporal_cloud_v1_replication_lag_p99 > 30",
        duracion="10m",
        caveats="Sin HA: no instalar. Percentil no agregable entre ns.",
        cardinalidad="Baja",
        requiere="Namespace HA",
        en_yaml="Sí (TemporalCloudReplicationLag)",
        dashboard="D02",
        pareja="—",
    ),
    # --- Limits ---
    dict(
        metrica="temporal_cloud_v1_operations_limit",
        categoria="Limits",
        tipo="Value",
        unidad="operations/s",
        desc_oficial="The current configured operations per second limit for a namespace.",
        desc_ops="Denominador. No alertar el gauge. Cambia si Temporal ajusta el tope (on-demand vs provisioned).",
        labels="— (solo comunes)",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Usar en ratio con operations_count. Alertar el ratio, no el limit.",
        promql="—",
        duracion="—",
        caveats="Gauge. Si el limit baja, el % sube sin más tráfico.",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="Sí (denominador)",
        dashboard="D03",
        pareja="temporal_cloud_v1_operations_count",
    ),
    dict(
        metrica="temporal_cloud_v1_action_limit",
        categoria="Limits",
        tipo="Value",
        unidad="actions/s",
        desc_oficial="The current configured actions per second limit for a namespace. Track utilization with total_action_count and is_background=false.",
        desc_ops="Denominador APS. Tablero. El paging es throttle, no este gauge.",
        labels="— (solo comunes; sin region — scope namespace)",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Ratio total_action / action_limit = P3. Throttle = P1.",
        promql="—",
        duracion="—",
        caveats="Sin label region. En provisioned, action_limit > envelope on-demand.",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="Sí (denominador ActionLimitHigh)",
        dashboard="D01, D03",
        pareja="temporal_cloud_v1_total_action_count, temporal_cloud_v1_action_on_demand_envelope_limit",
    ),
    dict(
        metrica="temporal_cloud_v1_service_request_limit",
        categoria="Limits",
        tipo="Value",
        unidad="requests/s",
        desc_oficial="The current configured frontend service RPS limit for a namespace. Track utilization with service_request_count.",
        desc_ops="Denominador RPS Frontend. P3 opcional: request_count / limit > 80%. El P1 es service_request_throttled.",
        labels="— (solo comunes)",
        optin="—",
        importancia="P3",
        orden=3,
        candidata="Sí",
        severidad="P3",
        condicion="Como denominador: service_request_count / service_request_limit > 0.8. No alertar el gauge solo.",
        promql="sum by (temporal_namespace) (temporal_cloud_v1_service_request_count) / sum by (temporal_namespace) (temporal_cloud_v1_service_request_limit) > 0.8",
        duracion="15m",
        caveats="Suma de operations puede no ser comparable 1:1 al limit (el limit es del frontend agregado). Validar en el ns.",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="No (candidato nuevo, simétrico a ActionLimitHigh)",
        dashboard="—",
        pareja="temporal_cloud_v1_service_request_count, temporal_cloud_v1_service_request_throttled_count",
    ),
    dict(
        metrica="temporal_cloud_v1_poller_limit",
        categoria="Limits",
        tipo="Value",
        unidad="pollers",
        desc_oficial="The current configured poller limit for a namespace. Track utilization with service_pending_requests.",
        desc_ops="Denominador. Alertar el ratio, no el limit.",
        labels="— (solo comunes)",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="Ver service_pending_requests (P3).",
        promql="—",
        duracion="—",
        caveats="—",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="Sí (denominador PollerLimitHigh)",
        dashboard="D03",
        pareja="temporal_cloud_v1_service_pending_requests",
    ),
    dict(
        metrica="temporal_cloud_v1_action_on_demand_envelope_limit",
        categoria="Limits",
        tipo="Value",
        unidad="actions/s",
        desc_oficial="On-demand envelope for actions/s. In provisioned mode: what the action limit would be on-demand. In on-demand: same as action_limit.",
        desc_ops="Diagnóstico TRU vs on-demand. Si action_limit ≠ envelope → provisioned. No alertar 0 ni igualdad.",
        labels="— (sin region; scope namespace)",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta. Comparar con action_limit para saber el modo.",
        promql="—",
        duracion="—",
        caveats="No incluye region.",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="No",
        dashboard="D03",
        pareja="temporal_cloud_v1_action_limit, temporal_cloud_v1_provisioned_capacity_tru_count",
    ),
    dict(
        metrica="temporal_cloud_v1_operations_on_demand_envelope_limit",
        categoria="Limits",
        tipo="Value",
        unidad="operations/s",
        desc_oficial="On-demand envelope for operations/s. Same semantics as action envelope.",
        desc_ops="Tablero de modo de capacidad. No alerta.",
        labels="— (sin region; scope namespace)",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta.",
        promql="—",
        duracion="—",
        caveats="No incluye region.",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_operations_limit",
    ),
    dict(
        metrica="temporal_cloud_v1_service_request_on_demand_envelope_limit",
        categoria="Limits",
        tipo="Value",
        unidad="requests/s",
        desc_oficial="On-demand envelope for service requests/s. Same semantics as action envelope.",
        desc_ops="Tablero de modo de capacidad. No alerta.",
        labels="— (sin region; scope namespace)",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta.",
        promql="—",
        duracion="—",
        caveats="No incluye region.",
        cardinalidad="Baja",
        requiere="—",
        en_yaml="No",
        dashboard="—",
        pareja="temporal_cloud_v1_service_request_limit",
    ),
    dict(
        metrica="temporal_cloud_v1_provisioned_capacity_tru_count",
        categoria="Limits",
        tipo="Value",
        unidad="TRU",
        desc_oficial="The number of Temporal Resource Units (TRUs) provisioned for a Namespace. 0 for Namespaces not in provisioned capacity mode.",
        desc_ops="0 = on-demand (normal). No alertar 0. Cada TRU ocioso puede facturar ns_capacity:tru (ver billable action_type).",
        labels="— (sin region; scope namespace)",
        optin="—",
        importancia="Dashboard",
        orden=4,
        candidata="No",
        severidad="—",
        condicion="No alerta. FinOps: cruzar con billable ns_capacity:tru si TRU > 0 y actions bajas.",
        promql="—",
        duracion="—",
        caveats="0 es el caso común. Alertar 0 genera ruido en todos los ns on-demand.",
        cardinalidad="Baja",
        requiere="Provisioned capacity",
        en_yaml="No",
        dashboard="D03",
        pareja="temporal_cloud_v1_billable_action_count",
    ),
]


HEADERS = [
    "ID",
    "Métrica",
    "Categoría",
    "Tipo OpenMetrics",
    "Unidad",
    "Descripción oficial (EN)",
    "Uso operativo / alerta",
    "Labels específicas",
    "Labels opt-in",
    "Importancia",
    "Orden",
    "Candidata alerta",
    "Severidad sugerida",
    "Condición de alerta",
    "PromQL borrador (v1, sin rate())",
    "for",
    "Caveats oficiales / notas",
    "Cardinalidad",
    "Requiere",
    "Ya en alertas-catalogo.yaml",
    "Dashboard BCP",
    "Métrica pareja",
    "Fuente",
]


def _styles():
    thin = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )
    return {
        "title": Font(name="Calibri", size=16, bold=True, color="FFFFFF"),
        "subtitle": Font(name="Calibri", size=11, color="1D2939"),
        "header": Font(name="Calibri", size=10, bold=True, color="FFFFFF"),
        "cell": Font(name="Calibri", size=10, color="1D2939"),
        "mono": Font(name="Consolas", size=9, color="1D2939"),
        "thin": thin,
        "fill_title": PatternFill("solid", fgColor="1B4F72"),
        "fill_header": PatternFill("solid", fgColor="154360"),
        "fill_p1": PatternFill("solid", fgColor="F97066"),
        "fill_p2": PatternFill("solid", fgColor="FDB022"),
        "fill_p3": PatternFill("solid", fgColor="FEC84B"),
        "fill_dash": PatternFill("solid", fgColor="B2DDFF"),
        "fill_none": PatternFill("solid", fgColor="D0D5DD"),
        "fill_yes": PatternFill("solid", fgColor="AAF0C4"),
        "fill_alt": PatternFill("solid", fgColor="F8FAFC"),
        "fill_white": PatternFill("solid", fgColor="FFFFFF"),
        "wrap": Alignment(wrap_text=True, vertical="top"),
        "wrap_c": Alignment(wrap_text=True, vertical="center", horizontal="center"),
    }


def _importance_fill(s, importancia):
    return {
        "P1": s["fill_p1"],
        "P2": s["fill_p2"],
        "P3": s["fill_p3"],
        "Dashboard": s["fill_dash"],
        "No alerta": s["fill_none"],
    }.get(importancia, s["fill_white"])


def _write_header_row(ws, headers, s, fill=None):
    fill = fill or s["fill_header"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(1, col, h)
        c.font = s["header"]
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = s["thin"]
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _apply_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _row_values(i, m):
    return [
        i,
        m["metrica"],
        m["categoria"],
        m["tipo"],
        m["unidad"],
        m["desc_oficial"],
        m["desc_ops"],
        m["labels"],
        m["optin"],
        m["importancia"],
        m["orden"],
        m["candidata"],
        m["severidad"],
        m["condicion"],
        m["promql"],
        m["duracion"],
        m["caveats"],
        m["cardinalidad"],
        m["requiere"],
        m["en_yaml"],
        m["dashboard"],
        m["pareja"],
        FUENTE,
    ]


def sheet_leyenda(wb, s):
    ws = wb.active
    ws.title = "Leyenda"
    ws.sheet_properties.tabColor = "1B4F72"

    ws.merge_cells("A1:F1")
    t = ws["A1"]
    t.value = "Temporal Cloud OpenMetrics — catálogo para candidatas a alerta"
    t.font = s["title"]
    t.fill = s["fill_title"]
    t.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28

    blurb = [
        ("Propósito", "Clasificar cada métrica temporal_cloud_v1_* por importancia. P1–P3 = candidatas a alerta. Dashboard / No alerta = tablero o denominador."),
        ("Fuente", f"{FUENTE}  |  corte {CORTE}"),
        ("Endpoint", f"{ENDPOINT}  (Bearer API key, rol Metrics Read-Only). Descriptores: {DESCRIPTORES}"),
        ("Convención v1", "Todas se exponen como gauge OpenMetrics. Rate = promedio /s en ventana de 1 min. Value = último agregado. Latency = percentil precalculado en segundos. NO usar rate()."),
        ("Labels comunes", COMMON_LABELS),
        ("Scrape", "Honor timestamps. Offset ~3 min. Scrape 30s. No habilitar labels opt-in (activity_type, nexus_*, worker_build_id) 24×7 salvo necesidad."),
        ("PromQL v0", "temporal_cloud_v0_* apagado 5 oct 2026. No copiar alertas con rate() del mixin legado."),
        ("Umbrales", "Los PromQL son borrador. Ajustar namespace, task queue, workflow type y umbrales por ambiente. No copiar a prod sin validar volumen."),
        ("Capa", "Estas métricas responden: ¿Cloud acepta y despacha trabajo? Worker saturado / lógica de activity → SDK metrics, no este catálogo."),
    ]
    ws["A3"] = "Campo"
    ws["B3"] = "Detalle"
    ws["A3"].font = s["header"]
    ws["B3"].font = s["header"]
    ws["A3"].fill = s["fill_header"]
    ws["B3"].fill = s["fill_header"]
    for r, (k, v) in enumerate(blurb, 4):
        ws.cell(r, 1, k).font = Font(name="Calibri", size=10, bold=True)
        ws.cell(r, 2, v).font = s["cell"]
        ws.cell(r, 1).alignment = s["wrap"]
        ws.cell(r, 2).alignment = s["wrap"]
        ws.cell(r, 1).border = s["thin"]
        ws.cell(r, 2).border = s["thin"]
        ws.row_dimensions[r].height = 36

    ws["A14"] = "Importancia"
    ws["B14"] = "Candidata alerta"
    ws["C14"] = "Significado"
    ws["D14"] = "Ejemplo"
    for col in range(1, 5):
        c = ws.cell(14, col)
        c.font = s["header"]
        c.fill = s["fill_header"]
        c.alignment = s["wrap_c"]
        c.border = s["thin"]

    legend = [
        ("P1", "Sí", "Paging. Trabajo rechazado, sin pollers, o error rate de ingreso.", "no_poller_tasks, *throttled*, resource_exhausted, service_error Start/Signal"),
        ("P2", "Sí", "Degradación / SLO. Fail rate, backlog, schedules perdidos, lag HA, silencio de Start.", "workflow_failed, backlog_age, schedule_missed_catchup, replication_lag_p99"),
        ("P3", "Sí", "Capacidad o FinOps. Ticket, no paging. % vs limit o gasto vs baseline.", "total_action / action_limit, billable vs 7d"),
        ("Dashboard", "No", "Útil en tablero o como denominador. No dispara incidente sola.", "workflow_success, poll_timeout, action_limit, p50"),
        ("No alerta", "No", "Deprecada, ruidosa o inaplicable por defecto.", "schedule_rate_limited_count (DEPRECATED)"),
    ]
    fills = [s["fill_p1"], s["fill_p2"], s["fill_p3"], s["fill_dash"], s["fill_none"]]
    for i, row in enumerate(legend):
        r = 15 + i
        for col, val in enumerate(row, 1):
            c = ws.cell(r, col, val)
            c.font = s["cell"]
            c.alignment = s["wrap"]
            c.border = s["thin"]
            if col == 1:
                c.fill = fills[i]
                c.alignment = s["wrap_c"]
                c.font = Font(name="Calibri", size=10, bold=True)
            elif col == 2:
                c.fill = s["fill_yes"] if val == "Sí" else s["fill_none"]
                c.alignment = s["wrap_c"]
        ws.row_dimensions[r].height = 40

    ws["A21"] = "Cómo leer las hojas"
    ws["A21"].font = Font(name="Calibri", size=12, bold=True, color="1B4F72")
    how = [
        "Catalogo: las 80 series oficiales. Filtro en la fila 1. Ordenar por columna Orden (1 = más importante) o Importancia.",
        "Candidatas alerta: solo P1–P3, ya filtradas. Empezar por P1. Las marcadas 'candidato nuevo' no están aún en grafana-temporal/alertas-catalogo.yaml.",
        "Resumen: conteo por importancia y por categoría.",
        "Convenciones: reglas oficiales (percentiles, backlog aproximado, opt-in, Nexus Preview, Standalone placeholders).",
        "Fuentes: URLs Temporal docs. No hay métricas inventadas fuera del metrics-reference.",
        "Scrape down (up{job=temporal-cloud}==0) no es una métrica Cloud: es de Prometheus. Sigue siendo P1 de ceguera; está en el YAML, no en este catálogo.",
    ]
    for i, line in enumerate(how):
        ws.cell(22 + i, 1, line).alignment = s["wrap"]
        ws.merge_cells(start_row=22 + i, start_column=1, end_row=22 + i, end_column=4)
        ws.row_dimensions[22 + i].height = 28

    _apply_widths(ws, [22, 22, 72, 72, 20, 20])
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_title_rows = "1:1"
    ws.oddHeader.left.text = "Temporal Cloud OpenMetrics — Leyenda"
    return ws


def sheet_catalogo(wb, s, rows, title, tab_color, filter_fn=None):
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = tab_color
    _write_header_row(ws, HEADERS, s)

    data = [m for m in rows if filter_fn(m)] if filter_fn else list(rows)
    data.sort(key=lambda m: (m["orden"], m["categoria"], m["metrica"]))

    for i, m in enumerate(data, 1):
        vals = _row_values(i, m)
        r = i + 1
        for col, val in enumerate(vals, 1):
            c = ws.cell(r, col, val)
            c.border = s["thin"]
            c.alignment = s["wrap"]
            if col in (2, 15):
                c.font = s["mono"]
            else:
                c.font = s["cell"]
            if col == 10:
                c.fill = _importance_fill(s, m["importancia"])
                c.alignment = s["wrap_c"]
                c.font = Font(name="Calibri", size=10, bold=True)
            elif col == 12:
                c.fill = s["fill_yes"] if m["candidata"] == "Sí" else s["fill_none"]
                c.alignment = s["wrap_c"]
            elif r % 2 == 0 and col not in (10, 12):
                c.fill = s["fill_alt"]
        ws.row_dimensions[r].height = 72

    last = len(data) + 1
    last_col = get_column_letter(len(HEADERS))
    ws.auto_filter.ref = f"A1:{last_col}{last}"

    widths = [6, 48, 24, 16, 18, 42, 42, 36, 28, 14, 8, 16, 28, 40, 48, 10, 42, 16, 24, 28, 16, 36, 22]
    _apply_widths(ws, widths)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    ws.oddHeader.left.text = f"Temporal Cloud OpenMetrics — {title}"
    ws.oddFooter.right.text = "Página &P de &N"
    return ws, len(data)


def sheet_resumen(wb, s, metrics):
    ws = wb.create_sheet("Resumen", 1)
    ws.sheet_properties.tabColor = "FDB022"

    ws.merge_cells("A1:D1")
    t = ws["A1"]
    t.value = "Conteo del catálogo oficial (80 series temporal_cloud_v1_*)"
    t.font = s["title"]
    t.fill = s["fill_title"]
    ws.row_dimensions[1].height = 28

    ws["A3"] = "Importancia"
    ws["B3"] = "Métricas"
    ws["C3"] = "Candidatas alerta"
    ws["D3"] = "% del catálogo"
    for col in range(1, 5):
        c = ws.cell(3, col)
        c.font = s["header"]
        c.fill = s["fill_header"]
        c.border = s["thin"]

    order = ["P1", "P2", "P3", "Dashboard", "No alerta"]
    total = len(metrics)
    counts = {k: sum(1 for m in metrics if m["importancia"] == k) for k in order}
    fills = [s["fill_p1"], s["fill_p2"], s["fill_p3"], s["fill_dash"], s["fill_none"]]
    for i, k in enumerate(order):
        r = 4 + i
        n = counts[k]
        cand = sum(1 for m in metrics if m["importancia"] == k and m["candidata"] == "Sí")
        ws.cell(r, 1, k).fill = fills[i]
        ws.cell(r, 1).font = Font(name="Calibri", size=10, bold=True)
        ws.cell(r, 2, n)
        ws.cell(r, 3, cand)
        ws.cell(r, 4, round(100 * n / total, 1))
        for col in range(1, 5):
            ws.cell(r, col).border = s["thin"]
            ws.cell(r, col).alignment = s["wrap_c"]

    ws.cell(9, 1, "Total")
    ws.cell(9, 2, total)
    ws.cell(9, 3, sum(1 for m in metrics if m["candidata"] == "Sí"))
    ws.cell(9, 4, 100)
    for col in range(1, 5):
        ws.cell(9, col).font = Font(name="Calibri", size=10, bold=True)
        ws.cell(9, col).border = s["thin"]
        ws.cell(9, col).fill = PatternFill("solid", fgColor="D5F5E3")

    chart = BarChart()
    chart.type = "col"
    chart.title = "Métricas por importancia"
    chart.y_axis.title = "Cantidad"
    data = Reference(ws, min_col=2, min_row=3, max_row=8)
    cats = Reference(ws, min_col=1, min_row=4, max_row=8)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend = None
    ws.add_chart(chart, "F3")

    ws["A12"] = "Categoría"
    ws["B12"] = "Métricas"
    ws["C12"] = "Candidatas (P1–P3)"
    for col in range(1, 4):
        c = ws.cell(12, col)
        c.font = s["header"]
        c.fill = s["fill_header"]
        c.border = s["thin"]

    cats_list = []
    for m in metrics:
        if m["categoria"] not in cats_list:
            cats_list.append(m["categoria"])
    for i, cat in enumerate(cats_list):
        r = 13 + i
        group = [m for m in metrics if m["categoria"] == cat]
        ws.cell(r, 1, cat)
        ws.cell(r, 2, len(group))
        ws.cell(r, 3, sum(1 for m in group if m["candidata"] == "Sí"))
        for col in range(1, 4):
            ws.cell(r, col).border = s["thin"]
            ws.cell(r, col).font = s["cell"]

    ws["A24"] = "Prioridad de implementación de alertas (recomendado)"
    ws["A24"].font = Font(name="Calibri", size=12, bold=True, color="1B4F72")
    steps = [
        "1. P1 ya en YAML: error rate Start/Signal, resource_exhausted, *throttled*, no_poller_tasks. Más scrape down (Prometheus up).",
        "2. P2 ya en YAML: fail rate workflow, backlog count, schedules (missed/buffer/overlap), open workflows jump, Start silence, replication_lag_p99 (solo HA).",
        "3. P2 candidatos nuevos: approximate_backlog_age (mejor que count), activity_fail/timeout/task_fail, service_latency_p95 con gate de volumen.",
        "4. P3: % vs action/operations/poller limit y billable vs baseline 7d (tickets FinOps, mute night).",
        "5. Nexus: no instalar hasta que el namespace use Nexus (Public Preview).",
        "6. Nunca alertar poll_timeout_count, TRU=0, ni percentiles en ns de poco tráfico sin gate.",
    ]
    for i, line in enumerate(steps):
        ws.cell(25 + i, 1, line).alignment = s["wrap"]
        ws.merge_cells(start_row=25 + i, start_column=1, end_row=25 + i, end_column=4)
        ws.row_dimensions[25 + i].height = 22

    _apply_widths(ws, [36, 14, 22, 16, 14, 18])
    return ws


def sheet_convenciones(wb, s):
    ws = wb.create_sheet("Convenciones")
    ws.sheet_properties.tabColor = "6941C6"
    ws.merge_cells("A1:B1")
    t = ws["A1"]
    t.value = "Convenciones oficiales OpenMetrics (no interpretaciones internas)"
    t.font = s["title"]
    t.fill = s["fill_title"]
    ws.row_dimensions[1].height = 28

    ws["A3"] = "Tema"
    ws["B3"] = "Qué dice Temporal"
    ws["A3"].font = s["header"]
    ws["B3"].font = s["header"]
    ws["A3"].fill = s["fill_header"]
    ws["B3"].fill = s["fill_header"]

    rows = [
        ("Tipos", "Rate = per-second rate of the 1-minute aggregate. Value = most recent aggregate (backlogs, limits). Percentile = pre-calculated latency in seconds. Exposed as OpenMetrics gauges."),
        ("Ventana 1 min", "A short spike can read below your limit even when it triggered throttling. Alert on *throttled* metrics, not only utilization %."),
        ("Percentiles y poco tráfico", "Need ~20 samples/window. Below that p50/p95/p99 converge on the slowest request. Gate alerts with service_request_count. Percentiles cannot be re-aggregated over longer windows."),
        ("No avg() de p95", "Avoid aggregating percentile metrics across dimensions (operation, workflow type, region): the percentile will not be accurate."),
        ("Labels comunes", COMMON_LABELS),
        ("Labels opt-in", "temporal_activity_type, temporal_worker_deployment_name, temporal_worker_build_id, nexus_service, nexus_operation. Query: /v1/metrics?labels=temporal_activity_type"),
        ("Actions sin region", "total_action_count, action_limit, envelopes y TRU no llevan label region (scope namespace)."),
        ("Backlog aproximado", "Overcount: invalid/expired tasks stay until they reach the head. Reset to 0 if the idle Task Queue is unloaded (~5 min without polls/new tasks/Describe). Started activities are not in the count."),
        ("Standalone Activities", "En métricas con temporal_workflow_type el placeholder es __standalone_activity. activity_terminate_count solo aplica a Standalone."),
        ("Nexus", "Public Preview. operation_* = caller namespace. request_* = handler namespace. Standalone Nexus: __temporal_standalone_nexus_operation__. Request latency de async start no espera el completion."),
        ("Billable vs cuota", "billable_action_count ≈ factura (±5%, con excepciones). total_action_count{is_background=false} vs action_limit = cuota APS. No son la misma cosa."),
        ("Resource exhausted vs throttle", "resource_exhausted = burst que un recurso no absorbe (SDK reintenta). *throttled* = rate limit de namespace. Pueden coexistir; el % vs limit puede estar verde."),
        ("Schedule DEPRECATED", "schedule_rate_limited_count se reemplaza por RPS estándar. No construir alerta nueva."),
        ("HA", "replication_lag_* solo namespaces high availability. Desactivar si no hay HA."),
        ("API", "GET /v1/metrics. Timestamp offset 3 min rounded down. honor_timestamps=true. GET /v1/descriptors para el catálogo programático."),
        ("SDK vs Cloud", "¿Cloud acepta trabajo? → este catálogo. ¿Worker saturado / activity lenta / dependency? → SDK metrics y traces."),
    ]
    for i, (k, v) in enumerate(rows):
        r = 4 + i
        ws.cell(r, 1, k).font = Font(name="Calibri", size=10, bold=True)
        ws.cell(r, 2, v).font = s["cell"]
        ws.cell(r, 1).alignment = s["wrap"]
        ws.cell(r, 2).alignment = s["wrap"]
        ws.cell(r, 1).border = s["thin"]
        ws.cell(r, 2).border = s["thin"]
        ws.row_dimensions[r].height = 48

    _apply_widths(ws, [28, 110])
    return ws


def sheet_fuentes(wb, s):
    ws = wb.create_sheet("Fuentes")
    ws.sheet_properties.tabColor = "027A48"
    ws.merge_cells("A1:C1")
    t = ws["A1"]
    t.value = "Fuentes oficiales Temporal"
    t.font = s["title"]
    t.fill = s["fill_title"]
    ws.row_dimensions[1].height = 28

    headers = ["Documento", "URL", "Para qué"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(3, col, h)
        c.font = s["header"]
        c.fill = s["fill_header"]
        c.border = s["thin"]

    sources = [
        ("OpenMetrics metrics reference", FUENTE, "Catálogo completo temporal_cloud_v1_* (este Excel)."),
        ("Set up Cloud metrics with OpenMetrics", "https://docs.temporal.io/cloud/metrics/openmetrics", "API key, rol Metrics Read-Only, scrape."),
        ("OpenMetrics API reference", "https://docs.temporal.io/cloud/metrics/openmetrics/api-reference", "Tipos Rate/Value, labels, honor_timestamps, offset 3 min, /v1/descriptors."),
        ("Set up metrics for Temporal Cloud", "https://docs.temporal.io/cloud/metrics", "Cuándo usar Cloud vs SDK."),
        ("OpenMetrics FAQ", "https://docs.temporal.io/cloud/metrics/openmetrics/faq", "Cardinalidad, series que faltan entre scrapes."),
        ("Migration guide v0 → v1", "https://docs.temporal.io/cloud/metrics/openmetrics/migration-guide", "No usar rate() en v1. PromQL v0 apagado 5 oct 2026."),
        ("Endpoint métricas", ENDPOINT, "Scrape Prometheus-compatible."),
        ("Endpoint descriptores", DESCRIPTORES, "Catálogo programático (misma API key)."),
        ("Standalone Activities", "https://docs.temporal.io/standalone-activity", "Placeholder __standalone_activity en labels."),
        ("Alertas BCP (interno)", "grafana-temporal/alertas-catalogo.yaml", "Reglas P1–P3 ya propuestas sobre este mismo set."),
    ]
    for i, (name, url, use) in enumerate(sources):
        r = 4 + i
        ws.cell(r, 1, name).font = s["cell"]
        cell_url = ws.cell(r, 2, url)
        cell_url.font = Font(name="Calibri", size=10, color="155EEF", underline="single")
        if url.startswith("http"):
            cell_url.hyperlink = url
        ws.cell(r, 3, use).font = s["cell"]
        for col in range(1, 4):
            ws.cell(r, col).border = s["thin"]
            ws.cell(r, col).alignment = s["wrap"]
        ws.row_dimensions[r].height = 28

    ws["A16"] = (
        f"Generado {CORTE}. {len(METRICS)} métricas. "
        "Si Temporal publica una serie nueva, volver a leer metrics-reference y regenerar con build_excel.py. "
        "No mezclar temporal_cloud_v0_*."
    )
    ws["A16"].alignment = s["wrap"]
    ws.merge_cells("A16:C16")
    ws.row_dimensions[16].height = 40

    _apply_widths(ws, [42, 88, 56])
    return ws


def main():
    assert len(METRICS) == 80, f"catálogo incompleto: {len(METRICS)} (esperado 80)"
    names = [m["metrica"] for m in METRICS]
    assert len(names) == len(set(names)), "nombres duplicados"

    wb = Workbook()
    s = _styles()
    sheet_leyenda(wb, s)
    sheet_resumen(wb, s, METRICS)
    sheet_catalogo(wb, s, METRICS, "Catalogo", "154360")
    sheet_catalogo(
        wb,
        s,
        METRICS,
        "Candidatas alerta",
        "B42318",
        filter_fn=lambda m: m["candidata"] == "Sí",
    )
    sheet_convenciones(wb, s)
    sheet_fuentes(wb, s)

    # Conditional formatting on Catalogo importancia column J
    cat = wb["Catalogo"]
    last = len(METRICS) + 1
    cat.conditional_formatting.add(
        f"J2:J{last}",
        FormulaRule(formula=['J2="P1"'], fill=s["fill_p1"]),
    )

    out = "metricas-temporal-cloud-openmetrics.xlsx"
    wb.save(out)
    candidatas = sum(1 for m in METRICS if m["candidata"] == "Sí")
    print(f"wrote {out}: {len(METRICS)} metrics, {candidatas} alert candidates")


if __name__ == "__main__":
    main()
