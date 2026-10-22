#!/usr/bin/env python3
"""Genera calculadora-privatelink.xlsx con fórmulas nativas de Excel."""

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

REGIONS = {
    "AWS": [
        ("us-east-1", "N. Virginia", False, False, 1.0),
        ("us-east-2", "Ohio", False, False, 1.0),
        ("us-west-2", "Oregon", False, False, 1.0),
        ("ca-central-1", "Canada Central", False, False, 1.0),
        ("ca-west-1", "Calgary", False, False, 1.0),
        ("sa-east-1", "São Paulo", True, False, 1.2),
        ("eu-central-1", "Frankfurt", True, False, 1.0),
        ("eu-central-2", "Zurich", False, False, 1.1),
        ("eu-west-1", "Ireland", False, False, 1.0),
        ("eu-west-2", "London", True, False, 1.1),
        ("eu-west-3", "Paris", False, False, 1.0),
        ("eu-north-1", "Stockholm", False, False, 1.0),
        ("eu-south-1", "Milan", False, False, 1.0),
        ("eu-south-2", "Spain", False, False, 1.0),
        ("ap-southeast-1", "Singapore", False, False, 1.1),
        ("ap-southeast-2", "Sydney", True, False, 1.1),
        ("ap-southeast-3", "Jakarta", False, False, 1.1),
        ("ap-northeast-1", "Tokyo", True, False, 1.1),
        ("ap-northeast-2", "Seoul", True, False, 1.1),
        ("ap-northeast-3", "Osaka", False, False, 1.1),
        ("ap-south-1", "Mumbai", True, False, 1.2),
        ("ap-south-2", "Hyderabad", False, False, 1.2),
        ("ap-east-1", "Hong Kong", False, False, 1.2),
        ("me-south-1", "Bahrain", False, False, 1.2),
        ("me-central-1", "UAE (no new customers)", True, False, 1.2),
        ("af-south-1", "Cape Town", True, False, 1.2),
        ("il-central-1", "Tel Aviv", False, False, 1.1),
    ],
    "Azure": [
        ("eastus", "Virginia", False, False, 1.0),
        ("eastus2", "Virginia", False, False, 1.0),
        ("centralus", "Iowa", False, False, 1.0),
        ("southcentralus", "Texas", False, False, 1.0),
        ("westus2", "Washington", False, False, 1.0),
        ("westus3", "Phoenix", False, False, 1.0),
        ("canadacentral", "Toronto", False, False, 1.0),
        ("brazilsouth", "São Paulo", True, False, 1.2),
        ("mexicocentral", "Mexico", True, False, 1.2),
        ("northeurope", "Ireland", False, False, 1.0),
        ("westeurope", "Netherlands", False, False, 1.0),
        ("uksouth", "London", True, False, 1.1),
        ("francecentral", "Paris", True, False, 1.1),
        ("germanywestcentral", "Frankfurt", True, False, 1.0),
        ("switzerlandnorth", "Zurich", True, False, 1.1),
        ("norwayeast", "Oslo", True, False, 1.1),
        ("swedencentral", "Sweden", False, False, 1.0),
        ("spaincentral", "Spain", False, False, 1.0),
        ("italynorth", "Milan", False, False, 1.0),
        ("austriaeast", "Vienna", False, False, 1.1),
        ("australiaeast", "NSW", True, False, 1.1),
        ("southeastasia", "Singapore", False, False, 1.1),
        ("eastasia", "Hong Kong", False, False, 1.2),
        ("japaneast", "Tokyo", True, False, 1.1),
        ("koreacentral", "Seoul", False, False, 1.1),
        ("centralindia", "Pune", True, False, 1.2),
        ("newzealandnorth", "New Zealand North", True, False, 1.1),
        ("uaenorth", "Dubai", True, False, 1.2),
        ("qatarcentral", "Doha", False, False, 1.2),
        ("southafricanorth", "Johannesburg", True, False, 1.2),
        ("jioindiawest", "Jio (sin Private Link)", False, True, 1.2),
        ("jioindiacentral", "Jio (sin Private Link)", False, True, 1.2),
    ],
    "GCP": [
        ("us-central1", "Iowa", False, False, 1.0),
        ("us-east1", "South Carolina", False, False, 1.0),
        ("us-east4", "N. Virginia", False, False, 1.0),
        ("us-west1", "Oregon", False, False, 1.0),
        ("us-west2", "Los Angeles", False, False, 1.0),
        ("us-west4", "Las Vegas", False, False, 1.0),
        ("us-south1", "Dallas", False, False, 1.0),
        ("northamerica-northeast1", "Montreal", False, False, 1.0),
        ("northamerica-northeast2", "Toronto", False, False, 1.0),
        ("southamerica-east1", "São Paulo", True, False, 1.2),
        ("southamerica-west1", "Santiago", True, False, 1.2),
        ("europe-west1", "Belgium", False, False, 1.0),
        ("europe-west2", "London", True, False, 1.1),
        ("europe-west3", "Frankfurt", False, False, 1.0),
        ("europe-west4", "Netherlands", False, False, 1.0),
        ("europe-west6", "Zurich", False, False, 1.1),
        ("europe-west8", "Milan", False, False, 1.0),
        ("europe-north1", "Finland", False, False, 1.0),
        ("europe-southwest1", "Madrid", False, False, 1.0),
        ("asia-east1", "Taiwan", False, False, 1.1),
        ("asia-east2", "Hong Kong (docs 1.2x)", False, False, 1.2),
        ("asia-northeast1", "Tokyo", False, False, 1.1),
        ("asia-south1", "Mumbai", True, False, 1.2),
        ("asia-south2", "Delhi", False, False, 1.2),
        ("asia-southeast1", "Singapore", False, False, 1.1),
        ("asia-southeast2", "Jakarta", False, False, 1.1),
        ("australia-southeast1", "Sydney", True, False, 1.1),
        ("australia-southeast2", "Melbourne", False, False, 1.1),
        ("me-west1", "Tel Aviv", False, False, 1.1),
        ("me-central2", "Dammam", False, False, 1.2),
    ],
}

TEAL = "00C4CC"
NAVY = "0F1419"
PANEL = "171E27"
YELLOW = "FFF3BF"
TEAL_FILL = "D4F4F5"
GREEN = "D3F3E2"
RED = "F8D4D2"
AMBER = "FCE8C3"
GRAY = "F4F6F8"
WHITE = "FFFFFF"
LINE = "C5CDD6"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
title_font = Font(name="Calibri", size=18, bold=True, color=NAVY)
h_font = Font(name="Calibri", size=12, bold=True, color=WHITE)
label_font = Font(name="Calibri", size=10, color="4A5560")
input_font = Font(name="Calibri", size=11, bold=True, color=NAVY)
formula_font = Font(name="Calibri", size=11, color=NAVY)
money_format = '"$"#,##0.00'
num_format = "#,##0.00"
int_format = "#,##0"


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def style_header_bar(ws, row, start, end, color=NAVY):
    for col in range(start, end + 1):
        cell = ws.cell(row, col)
        cell.fill = fill(color)
        cell.font = h_font
        cell.alignment = Alignment(vertical="center")


def input_cell(ws, row, col, value, number_format=None):
    cell = ws.cell(row, col, value)
    cell.fill = fill(YELLOW)
    cell.font = input_font
    cell.border = thin
    cell.alignment = Alignment(horizontal="left", vertical="center")
    if number_format:
        cell.number_format = number_format
    return cell


def formula_cell(ws, row, col, formula, number_format=None, money=False):
    if isinstance(formula, str) and formula and not formula.startswith("="):
        formula = "=" + formula
    cell = ws.cell(row, col, formula)
    cell.font = formula_font
    cell.border = thin
    cell.fill = fill(GRAY)
    if money or number_format == money_format:
        cell.number_format = money_format
        cell.alignment = Alignment(horizontal="right", vertical="center")
    elif number_format:
        cell.number_format = number_format
    return cell


def label(ws, row, col, text):
    cell = ws.cell(row, col, text)
    cell.font = label_font
    cell.alignment = Alignment(vertical="center")
    return cell


def build():
    wb = Workbook()

    # --- Regiones ---
    wr = wb.active
    wr.title = "Regiones"
    headers = ["Cloud", "Region", "Ubicacion", "Premium_cross_cloud", "Jio_sin_PL", "Multiplier", "Clave"]
    for i, h in enumerate(headers, 1):
        wr.cell(1, i, h)
    style_header_bar(wr, 1, 1, 7)
    wr.auto_filter.ref = "A1:G200"
    wr.freeze_panes = "A2"

    row = 2
    ranges = {}
    for cloud, items in REGIONS.items():
        start = row
        for rid, loc, prem, jio, mult in items:
            wr.cell(row, 1, cloud)
            wr.cell(row, 2, rid)
            wr.cell(row, 3, loc)
            wr.cell(row, 4, "SI" if prem else "NO")
            wr.cell(row, 5, "SI" if jio else "NO")
            wr.cell(row, 6, mult)
            wr.cell(row, 6).number_format = "0.0"
            wr.cell(row, 7, f"{cloud}|{rid}")
            if prem:
                wr.cell(row, 4).fill = fill(AMBER)
            if jio:
                wr.cell(row, 5).fill = fill(RED)
            row += 1
        ranges[cloud] = (start, row - 1)

    last = row - 1
    for col, width in zip(range(1, 8), [12, 28, 28, 22, 14, 12, 36]):
        wr.column_dimensions[get_column_letter(col)].width = width
    wr.sheet_properties.tabColor = "8B9AAB"

    for cloud, (a, b) in ranges.items():
        wb.defined_names.add(DefinedName(f"Cloud_{cloud}", attr_text=f"Regiones!$B${a}:$B${b}"))
    wb.defined_names.add(DefinedName("AllClaves", attr_text=f"Regiones!$G$2:$G${last}"))
    wb.defined_names.add(DefinedName("AllClouds", attr_text="Listas!$A$2:$A$4"))

    # --- Listas ---
    wl = wb.create_sheet("Listas")
    wl["A1"] = "Cloud"
    wl["A2"] = "AWS"
    wl["A3"] = "Azure"
    wl["A4"] = "GCP"
    wl["B1"] = "Availability"
    wl["B2"] = "Multi-zone"
    wl["B3"] = "Single-zone"
    wl["C1"] = "Direccion"
    wl["C2"] = "Unidirectional"
    wl["C3"] = "Bidirectional"
    wl["C4"] = "None"
    style_header_bar(wl, 1, 1, 3)
    for col, width in zip("ABC", [14, 16, 26]):
        wl.column_dimensions[col].width = width
    wl.sheet_state = "visible"

    # --- Tarifas ---
    wt = wb.create_sheet("Tarifas", 0)
    wt["A1"] = "Tarifas y constantes"
    wt["A1"].font = title_font
    wt.merge_cells("A1:D1")
    wt["A2"] = "Amarillo = editable. Cluster Linking hourly/GB son list prices oficiales. Dedicated CKU/ingress/egress/storage no son list price público."
    wt["A2"].font = Font(name="Calibri", size=10, italic=True, color="4A5560")
    wt.merge_cells("A2:D2")

    wt["A4"] = "Cluster Linking (oficial, Billing Dimensions)"
    style_header_bar(wt, 4, 1, 3)
    wt.merge_cells("A4:C4")
    official = [
        (5, "CL_PUBLIC_HOUR", 0.25, "USD/h por cluster si el endpoint es internet"),
        (6, "CL_PRIVATE_HOUR", 0.50, "USD/h por cluster intra-cloud private"),
        (7, "CL_XCLOUD_HOUR", 2.50, "USD/h por cluster cross-cloud private"),
        (8, "CL_XCLOUD_PREMIUM_HOUR", 5.00, "USD/h si esa región es Premium"),
        (9, "CL_READ_GB", 0.25, "USD/GB ClusterLinkingRead (incluye metadata)"),
    ]
    wt["A5"] = "Clave"
    wt["B5"] = "USD"
    wt["C5"] = "Nota"
    # overwrite row 5 with first official after header - wait I used row 4 as title
    # Put header on 5, data 6+
    wt["A4"] = "Cluster Linking (oficial, Billing Dimensions)"
    wt["A5"] = "Clave"
    wt["B5"] = "USD"
    wt["C5"] = "Nota"
    style_header_bar(wt, 5, 1, 3)
    official = [
        (6, "CL_PUBLIC_HOUR", 0.25, "USD/h por cluster si el endpoint es internet"),
        (7, "CL_PRIVATE_HOUR", 0.50, "USD/h por cluster intra-cloud private"),
        (8, "CL_XCLOUD_HOUR", 2.50, "USD/h por cluster cross-cloud private"),
        (9, "CL_XCLOUD_PREMIUM_HOUR", 5.00, "USD/h si esa región es Premium (docs Regions)"),
        (10, "CL_READ_GB", 0.25, "USD/GB ClusterLinkingRead. Confluent paga egress CSP"),
    ]
    names = {}
    for r, key, val, note in official:
        wt.cell(r, 1, key).font = Font(name="Calibri", size=10)
        input_cell(wt, r, 2, val, money_format)
        wt.cell(r, 3, note).font = Font(name="Calibri", size=9, color="4A5560")
        names[key] = f"Tarifas!$B${r}"
        wb.defined_names.add(DefinedName(key, attr_text=f"Tarifas!$B${r}"))

    wt["A12"] = "Dedicated (editable — no es list price público)"
    style_header_bar(wt, 12, 1, 3)
    wt.merge_cells("A12:C12")
    wt["A13"] = "Clave"
    wt["B13"] = "USD"
    wt["C13"] = "Nota"
    style_header_bar(wt, 13, 1, 3)
    dedicated = [
        (14, "CKU_HOUR", 1.50, "USD/CKU-hora base. Sacar de: confluent billing price list"),
        (15, "KAFKA_IN_GB", 0.10, "KafkaNetworkWrite USD/GB base"),
        (16, "KAFKA_OUT_GB", 0.05, "KafkaNetworkRead USD/GB base"),
        (17, "STORAGE_GB_MONTH", 0.08, "USD/GB-mes post-replication"),
        (18, "CSP_PE_HOUR", 0.01, "Azure PE / AWS VPCE-AZ / GCP PSC USD/hora (list US)"),
        (19, "CSP_PE_GB", 0.01, "CSP data processed USD/GB (list US)"),
        (20, "HOURS_MONTH", 730, "Horas / mes (24×365/12)"),
        (21, "RF", 3, "Replication factor Dedicated. Storage ≈ 3×"),
        (22, "CKU_IN_MBPS", 60, "Ingress MBps por CKU (docs cluster types)"),
        (23, "CKU_OUT_MBPS", 180, "Egress MBps por CKU"),
        (24, "CKU_PARTITIONS", 4500, "Partitions pre-replication por CKU"),
    ]
    for r, key, val, note in dedicated:
        wt.cell(r, 1, key).font = Font(name="Calibri", size=10)
        fmt = money_format if r <= 19 else ("0" if r >= 21 else "0")
        if key == "HOURS_MONTH":
            fmt = "0"
        input_cell(wt, r, 2, val, fmt)
        wt.cell(r, 3, note).font = Font(name="Calibri", size=9, color="4A5560")
        wb.defined_names.add(DefinedName(key, attr_text=f"Tarifas!$B${r}"))

    wt["A26"] = "Tope CKU Dedicated por cloud"
    style_header_bar(wt, 26, 1, 2)
    wt["A27"] = "AWS"
    wt["B27"] = 252
    wt["A28"] = "Azure"
    wt["B28"] = 100
    wt["A29"] = "GCP"
    wt["B29"] = 152
    for r in range(27, 30):
        wt.cell(r, 2).number_format = "0"
        wt.cell(r, 2).border = thin

    wt["A31"] = "Fuentes oficiales"
    style_header_bar(wt, 31, 1, 2)
    sources = [
        "https://docs.confluent.io/cloud/current/billing/billing-dimensions.html",
        "https://docs.confluent.io/cloud/current/multi-cloud/cluster-linking/private-networking.html",
        "https://docs.confluent.io/cloud/current/get-started/regions.html",
        "https://docs.confluent.io/cloud/current/networking/private-links/azure-privatelink.html",
        "https://docs.confluent.io/cloud/current/networking/private-links/aws-privatelink.html",
        "https://docs.confluent.io/cloud/current/networking/private-links/gcp-private-service-connect.html",
        "https://docs.confluent.io/cloud/current/clusters/cluster-types.html",
    ]
    for i, url in enumerate(sources):
        wt.cell(32 + i, 1, url).hyperlink = url
        wt.cell(32 + i, 1).font = Font(name="Calibri", size=9, color="0563C1", underline="single")
        wt.merge_cells(start_row=32 + i, start_column=1, end_row=32 + i, end_column=3)

    wt.column_dimensions["A"].width = 28
    wt.column_dimensions["B"].width = 14
    wt.column_dimensions["C"].width = 78
    wt.row_dimensions[1].height = 24
    wt.sheet_properties.tabColor = TEAL

    # --- Calculadora ---
    ws = wb.create_sheet("Calculadora", 0)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_title_rows = "1:3"
    ws.oddHeader.left.text = "Confluent Cloud Dedicated · Private Link"
    ws.oddFooter.left.text = "USD · GiB · Dedicated only. Cluster Linking list prices oficiales."
    ws.oddFooter.right.text = "Página &P de &N"

    ws.merge_cells("A1:L1")
    ws["A1"] = "Calculadora Private Link · Confluent Cloud Dedicated"
    ws["A1"].font = title_font
    ws.merge_cells("A2:L2")
    ws["A2"] = (
        "Celdas amarillas = input. El resto son fórmulas. Cluster Linking $0.25 / $0.50 / $2.50 / $5.00 y $0.25/GB son oficiales. "
        "Dedicated CKU/ingress/egress/storage: editar en hoja Tarifas con `confluent billing price list`."
    )
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="4A5560")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 32

    # Cluster A
    ws.merge_cells("A4:C4")
    ws["A4"] = "Cluster A · source"
    style_header_bar(ws, 4, 1, 3, "0078D4")
    fields_a = [
        (5, "Cloud", "Azure"),
        (6, "Region", "eastus2"),
        (7, "Availability", "Multi-zone"),
        (8, "CKU", 2),
        (9, "VNets / VPCs cliente", 1),
        (10, "Produce cliente (GB/mes)", 400),
        (11, "Consume cliente (GB/mes)", 800),
        (12, "Storage pre-RF (GB)", 200),
    ]
    for r, name, val in fields_a:
        label(ws, r, 1, name)
        if isinstance(val, str):
            input_cell(ws, r, 2, val)
        else:
            input_cell(ws, r, 2, val, int_format)

    # Cluster B
    ws.merge_cells("E4:G4")
    ws["E4"] = "Cluster B · destination"
    style_header_bar(ws, 4, 5, 7, "FF9900")
    fields_b = [
        (5, "Cloud", "AWS"),
        (6, "Region", "us-east-1"),
        (7, "Availability", "Multi-zone"),
        (8, "CKU", 2),
        (9, "VNets / VPCs cliente", 1),
        (10, "Produce cliente (GB/mes)", 0),
        (11, "Consume cliente (GB/mes)", 400),
        (12, "Storage pre-RF (GB)", 200),
    ]
    for r, name, val in fields_b:
        label(ws, r, 5, name)
        if isinstance(val, str):
            input_cell(ws, r, 6, val)
        else:
            input_cell(ws, r, 6, val, int_format)

    # Linking
    ws.merge_cells("A14:C14")
    ws["A14"] = "Cluster Linking"
    style_header_bar(ws, 14, 1, 3, TEAL)
    label(ws, 15, 1, "Direccion")
    input_cell(ws, 15, 2, "Unidirectional")
    label(ws, 16, 1, "Mirror A → B (GB/mes)")
    input_cell(ws, 16, 2, 500, int_format)
    label(ws, 17, 1, "Mirror B → A (GB/mes)")
    input_cell(ws, 17, 2, 0, int_format)

    # Lookups / diagnosis
    ws.merge_cells("E14:G14")
    ws["E14"] = "Diagnóstico (fórmulas)"
    style_header_bar(ws, 14, 5, 7, PANEL)

    # Helper lookup: INDEX/MATCH on Cloud+Region
    # Premium A
    look_prem = lambda cloud, region: (
        f'IFERROR(INDEX(Regiones!$D$2:$D${last},MATCH({cloud}&"|"&{region},Regiones!$G$2:$G${last},0)),"NO")'
    )
    look_jio = lambda cloud, region: (
        f'IFERROR(INDEX(Regiones!$E$2:$E${last},MATCH({cloud}&"|"&{region},Regiones!$G$2:$G${last},0)),"NO")'
    )
    look_mult = lambda cloud, region: (
        f'IFERROR(INDEX(Regiones!$F$2:$F${last},MATCH({cloud}&"|"&{region},Regiones!$G$2:$G${last},0)),1)'
    )

    diags = [
        (15, "Premium A", look_prem("$B$5", "$B$6")),
        (16, "Jio A (sin PL)", look_jio("$B$5", "$B$6")),
        (17, "Multiplier A", look_mult("$B$5", "$B$6")),
        (18, "Premium B", look_prem("$F$5", "$F$6")),
        (19, "Jio B (sin PL)", look_jio("$F$5", "$F$6")),
        (20, "Multiplier B", look_mult("$F$5", "$F$6")),
        (21, "Cross-cloud", '=$B$5<>$F$5'),
        (22, "Links", '=IF(OR($B$15="None",$B$15="Sin Cluster Linking"),0,IF($B$15="Bidirectional",2,1))'),
        (23, "GB A→B efectivo", '=IF(OR($B$15="None",$B$15="Sin Cluster Linking"),0,$B$16)'),
        (24, "GB B→A efectivo", '=IF($B$15="Bidirectional",$B$17,0)'),
        (25, "Hourly A",
         '=IF($F$22=0,0,IF($F$21,IF($F$15="SI",CL_XCLOUD_PREMIUM_HOUR,CL_XCLOUD_HOUR),CL_PRIVATE_HOUR*$F$17))'),
        (26, "Hourly B",
         '=IF($F$22=0,0,IF($F$21,IF($F$18="SI",CL_XCLOUD_PREMIUM_HOUR,CL_XCLOUD_HOUR),CL_PRIVATE_HOUR*$F$20))'),
        (27, "AZ por red A", '=IF($B$7="Single-zone",1,3)'),
        (28, "AZ por red B", '=IF($F$7="Single-zone",1,3)'),
        (29, "Unidades CSP A", '=$B$9*$F$27'),
        (30, "Unidades CSP B", '=$F$9*$F$28'),
        (31, "Endpoints a crear A", '=IF($B$5="AWS",$B$9,$B$9*$F$27)'),
        (32, "Endpoints a crear B", '=IF($F$5="AWS",$F$9,$F$9*$F$28)'),
        (33, "Tipo endpoint A",
         '=IF($B$5="AWS","VPC endpoint (1 × N AZ)",IF($B$5="Azure","Private Endpoints","PSC endpoints"))'),
        (34, "Tipo endpoint B",
         '=IF($F$5="AWS","VPC endpoint (1 × N AZ)",IF($F$5="Azure","Private Endpoints","PSC endpoints"))'),
        (35, "Región válida A",
         f'=IF(COUNTIFS(Regiones!$A$2:$A${last},$B$5,Regiones!$B$2:$B${last},$B$6)=1,"OK","Region no coincide con cloud")'),
        (36, "Región válida B",
         f'=IF(COUNTIFS(Regiones!$A$2:$A${last},$F$5,Regiones!$B$2:$B${last},$F$6)=1,"OK","Region no coincide con cloud")'),
        (37, "Mismo Organization", "OBLIGATORIO si ambos son private Dedicated"),
    ]
    for r, name, formula in diags:
        label(ws, r, 5, name)
        formula_cell(ws, r, 6, formula)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
    ws["F17"].number_format = "0.0"
    ws["F20"].number_format = "0.0"
    ws["F25"].number_format = money_format
    ws["F26"].number_format = money_format
    for r in (22, 23, 24, 27, 28, 29, 30, 31, 32):
        ws[f"F{r}"].number_format = int_format

    # CKU suggested
    # mbps = GB * 1024 / (hours * 3600)
    ws["E38"] = "CKU sugerido A"
    ws["E38"].font = label_font
    formula_cell(
        ws, 38, 6,
        '=MAX(IF($B$7="Multi-zone",2,1),ROUNDUP(MAX((($B$10+$F$24)*1024)/(HOURS_MONTH*3600)/CKU_IN_MBPS,(($B$11+$F$23)*1024)/(HOURS_MONTH*3600)/CKU_OUT_MBPS),0))',
        int_format,
    )
    ws.merge_cells("F38:G38")
    ws["E39"] = "CKU sugerido B"
    ws["E39"].font = label_font
    formula_cell(
        ws, 39, 6,
        '=MAX(IF($F$7="Multi-zone",2,1),ROUNDUP(MAX((($F$10+$F$23)*1024)/(HOURS_MONTH*3600)/CKU_IN_MBPS,(($F$11+$F$24)*1024)/(HOURS_MONTH*3600)/CKU_OUT_MBPS),0))',
        int_format,
    )
    ws.merge_cells("F39:G39")
    ws["E40"] = "Alerta CKU / Jio / PL"
    ws["E40"].font = label_font
    formula_cell(
        ws, 40, 6,
        '=TRIM(SUBSTITUTE(SUBSTITUTE('
        'IF(OR($F$16="SI",$F$19="SI"),"Jio: no Private Link · ","")&'
        'IF(OR($B$8<$F$38,$F$8<$F$39),"CKU por debajo de la guía · ","")&'
        'IF(OR(AND($B$7="Multi-zone",$B$8<2),AND($F$7="Multi-zone",$F$8<2)),"Multi-zone exige 2 CKU · ","")&'
        'IF($F$21,"Cross-cloud private Cluster Linking · ","Intra-cloud private · ")&'
        'IF(OR($F$15="SI",$F$18="SI"),"Premium region $5.00/h · ","")&'
        'IF($F$22=2,"Bidirectional = 2 links · ","")&'
        'IF(OR($F$35<>"OK",$F$36<>"OK"),"Region no coincide con cloud","")'
        '," ·  "," · "), " ·", ""))',
    )
    ws.merge_cells("F40:G40")
    ws["F40"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[40].height = 32

    # Results
    ws.merge_cells("I4:L4")
    ws["I4"] = "Estimado mensual USD"
    style_header_bar(ws, 4, 9, 12, TEAL)
    ws["I5"] = "Concepto"
    ws["J5"] = "USD / mes"
    ws["K5"] = "Cómo se calcula"
    ws["L5"] = "Bloque"
    style_header_bar(ws, 5, 9, 12, PANEL)

    lines = [
        (6, "CKU cluster A",
         "=$B$8*CKU_HOUR*$F$17*HOURS_MONTH",
         '=$B$8&" CKU × "&TEXT(CKU_HOUR*$F$17,"$0.00")&"/h × "&HOURS_MONTH&"h"',
         "Confluent"),
        (7, "CKU cluster B",
         "=$F$8*CKU_HOUR*$F$20*HOURS_MONTH",
         '=$F$8&" CKU × "&TEXT(CKU_HOUR*$F$20,"$0.00")&"/h × "&HOURS_MONTH&"h"',
         "Confluent"),
        (8, "Kafka ingress clientes A",
         "=$B$10*KAFKA_IN_GB*$F$17",
         '=$B$10&" GB × "&TEXT(KAFKA_IN_GB*$F$17,"$0.00")',
         "Confluent"),
        (9, "Kafka ingress clientes B",
         "=$F$10*KAFKA_IN_GB*$F$20",
         '=$F$10&" GB × "&TEXT(KAFKA_IN_GB*$F$20,"$0.00")',
         "Confluent"),
        (10, "Kafka egress clientes A",
         "=$B$11*KAFKA_OUT_GB*$F$17",
         '=$B$11&" GB × "&TEXT(KAFKA_OUT_GB*$F$17,"$0.00")',
         "Confluent"),
        (11, "Kafka egress clientes B",
         "=$F$11*KAFKA_OUT_GB*$F$20",
         '=$F$11&" GB × "&TEXT(KAFKA_OUT_GB*$F$20,"$0.00")',
         "Confluent"),
        (12, "Storage A (× RF)",
         "=$B$12*RF*STORAGE_GB_MONTH",
         '=$B$12&" GB pre × "&RF&" × "&TEXT(STORAGE_GB_MONTH,"$0.00")',
         "Confluent"),
        (13, "Storage B (× RF)",
         "=$F$12*RF*STORAGE_GB_MONTH",
         '=$F$12&" GB pre × "&RF&" × "&TEXT(STORAGE_GB_MONTH,"$0.00")',
         "Confluent"),
        (14, "Cluster link hourly A",
         "=$F$25*$F$22*HOURS_MONTH",
         '=TEXT($F$25,"$0.00")&"/h × "&$F$22&" link(s) × "&HOURS_MONTH&"h"',
         "Confluent"),
        (15, "Cluster link hourly B",
         "=$F$26*$F$22*HOURS_MONTH",
         '=TEXT($F$26,"$0.00")&"/h × "&$F$22&" link(s) × "&HOURS_MONTH&"h"',
         "Confluent"),
        (16, "ClusterLinkingRead A→B",
         "=$F$23*CL_READ_GB*$F$17",
         '=$F$23&" GB × "&TEXT(CL_READ_GB,"$0.00")&" × "&$F$17',
         "Confluent"),
        (17, "ClusterLinkingWrite A→B",
         "=$F$23*KAFKA_IN_GB*$F$20",
         '=$F$23&" GB × dest KafkaNetworkWrite"',
         "Confluent"),
        (18, "ClusterLinkingRead B→A",
         "=$F$24*CL_READ_GB*$F$20",
         '=$F$24&" GB × "&TEXT(CL_READ_GB,"$0.00")&" × "&$F$20',
         "Confluent"),
        (19, "ClusterLinkingWrite B→A",
         "=$F$24*KAFKA_IN_GB*$F$17",
         '=$F$24&" GB × dest KafkaNetworkWrite"',
         "Confluent"),
        (20, "CSP endpoints A",
         "=$F$29*CSP_PE_HOUR*HOURS_MONTH+($B$10+$B$11)*CSP_PE_GB",
         '=$F$29&" unidades + "&($B$10+$B$11)&" GB cliente"',
         "CSP"),
        (21, "CSP endpoints B",
         "=$F$30*CSP_PE_HOUR*HOURS_MONTH+($F$10+$F$11)*CSP_PE_GB",
         '=$F$30&" unidades + "&($F$10+$F$11)&" GB cliente"',
         "CSP"),
    ]
    for r, name, usd, how, block in lines:
        ws.cell(r, 9, name).font = Font(name="Calibri", size=10)
        ws.cell(r, 9).border = thin
        formula_cell(ws, r, 10, usd, money=True)
        formula_cell(ws, r, 11, how)
        ws.cell(r, 12, block).border = thin
        ws.cell(r, 12).font = Font(name="Calibri", size=9, color="4A5560")
        if block == "CSP":
            ws.cell(r, 9).fill = fill("EEF6FF")
            ws.cell(r, 12).fill = fill("EEF6FF")

    ws.merge_cells("I22:I22")
    ws["I22"] = "Subtotal Confluent"
    ws["I22"].font = Font(name="Calibri", size=11, bold=True)
    formula_cell(ws, 22, 10, "=SUMIF($L$6:$L$21,\"Confluent\",$J$6:$J$21)", money=True)
    ws["J22"].fill = fill(TEAL_FILL)
    ws["J22"].font = Font(name="Calibri", size=11, bold=True)
    ws["I23"] = "Subtotal CSP"
    ws["I23"].font = Font(name="Calibri", size=11, bold=True)
    formula_cell(ws, 23, 10, "=SUMIF($L$6:$L$21,\"CSP\",$J$6:$J$21)", money=True)
    ws["I24"] = "TOTAL USD / mes"
    ws["I24"].font = Font(name="Calibri", size=13, bold=True, color=WHITE)
    ws["I24"].fill = fill(NAVY)
    ws["J24"] = "=J22+J23"
    ws["J24"].font = Font(name="Calibri", size=14, bold=True, color=WHITE)
    ws["J24"].fill = fill(NAVY)
    ws["J24"].number_format = money_format
    ws["J24"].border = thin
    ws["K24"] = "Confluent + endpoints del cloud provider"
    ws["K24"].font = Font(name="Calibri", size=9, italic=True, color="4A5560")

    ws["I26"] = "Private Link a provisionar"
    style_header_bar(ws, 26, 9, 12, PANEL)
    ws.merge_cells("I26:L26")
    ws["I27"] = "A"
    formula_cell(ws, 27, 10, '=$F$31&" × "&$F$33')
    ws.merge_cells("J27:L27")
    ws["I28"] = "B"
    formula_cell(ws, 28, 10, '=$F$32&" × "&$F$34')
    ws.merge_cells("J28:L28")
    ws["I29"] = "DNS por cluster"
    formula_cell(ws, 29, 10, '=IF($B$7="Single-zone","* + 1 zonal (A records)","* + 3 zonales (A/CNAME). No hardcodear brokers.")')
    ws.merge_cells("J29:L29")
    ws["I30"] = "CKU A / B vs guía"
    formula_cell(ws, 30, 10, '=$B$8&" / "&$F$8&"  (sugerido "&$F$38&" / "&$F$39&")")')
    ws.merge_cells("J30:L30")

    # Chart data (compact)
    ws["I32"] = "Gráfico"
    style_header_bar(ws, 32, 9, 10, PANEL)
    ws["I33"] = "CKU"
    ws["J33"] = "=J6+J7"
    ws["I34"] = "Kafka in/out + storage"
    ws["J34"] = "=SUM(J8:J13)"
    ws["I35"] = "Cluster Linking"
    ws["J35"] = "=SUM(J14:J19)"
    ws["I36"] = "CSP Private Link"
    ws["J36"] = "=J23"
    for r in range(33, 37):
        ws.cell(r, 9).border = thin
        ws.cell(r, 10).number_format = money_format
        ws.cell(r, 10).border = thin
        ws.cell(r, 10).fill = fill(GRAY)

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "USD / mes por bloque"
    chart.y_axis.title = None
    chart.x_axis.title = None
    chart.shape = 4
    data = Reference(ws, min_col=10, min_row=33, max_row=36)
    cats = Reference(ws, min_col=9, min_row=33, max_row=36)
    chart.add_data(data, from_rows=False, titles_from_data=False)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend = None
    chart.height = 7
    chart.width = 12
    if chart.series:
        chart.series[0].graphicalProperties.solidFill = TEAL
    ws.add_chart(chart, "I38")

    # Data validations
    dv_cloud = DataValidation(type="list", formula1="=AllClouds", allow_blank=False)
    dv_cloud.error = "Elige AWS, Azure o GCP"
    dv_cloud.errorTitle = "Cloud"
    dv_az = DataValidation(type="list", formula1="=Listas!$B$2:$B$3", allow_blank=False)
    dv_dir = DataValidation(type="list", formula1="=Listas!$C$2:$C$4", allow_blank=False)
    dv_reg_a = DataValidation(type="list", formula1='=INDIRECT("Cloud_"&$B$5)', allow_blank=False)
    dv_reg_a.error = "Elige una región del cloud de la celda B5. Cambia Cloud primero."
    dv_reg_a.errorTitle = "Region A"
    dv_reg_b = DataValidation(type="list", formula1='=INDIRECT("Cloud_"&$F$5)', allow_blank=False)
    dv_reg_b.error = "Elige una región del cloud de la celda F5. Cambia Cloud primero."
    dv_reg_b.errorTitle = "Region B"
    for dv in (dv_cloud, dv_az, dv_dir, dv_reg_a, dv_reg_b):
        dv.showErrorMessage = True
        dv.showInputMessage = True
        ws.add_data_validation(dv)
    dv_cloud.add("B5")
    dv_cloud.add("F5")
    dv_reg_a.add("B6")
    dv_reg_b.add("F6")
    dv_az.add("B7")
    dv_az.add("F7")
    dv_dir.add("B15")

    # Conditional formatting
    ws.conditional_formatting.add(
        "F35:G36",
        FormulaRule(formula=['$F$35<>"OK"'], fill=fill(RED)),
    )
    ws.conditional_formatting.add(
        "F16",
        FormulaRule(formula=['$F$16="SI"'], fill=fill(RED)),
    )
    ws.conditional_formatting.add(
        "F19",
        FormulaRule(formula=['$F$19="SI"'], fill=fill(RED)),
    )
    ws.conditional_formatting.add(
        "F15",
        FormulaRule(formula=['$F$15="SI"'], fill=fill(AMBER)),
    )
    ws.conditional_formatting.add(
        "F18",
        FormulaRule(formula=['$F$18="SI"'], fill=fill(AMBER)),
    )

    # Column widths / row heights
    widths = {
        "A": 26, "B": 22, "C": 8, "D": 3,
        "E": 22, "F": 18, "G": 16, "H": 3,
        "I": 28, "J": 16, "K": 48, "L": 12,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for r in range(5, 13):
        ws.row_dimensions[r].height = 18
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A4"

    # Notes on left bottom
    ws.merge_cells("A19:C19")
    ws["A19"] = "Cómo usar"
    style_header_bar(ws, 19, 1, 3, PANEL)
    ws.merge_cells("A20:C37")
    ws["A20"] = (
        "1. Elige Cloud y Region de A y B (el dropdown de Region depende del Cloud).\n"
        "2. Multi-zone Dedicated exige mínimo 2 CKU y 3 endpoints (Azure/GCP) o 1 VPCE en 3 AZ (AWS).\n"
        "3. Dirección: Unidirectional = 1 link; Bidirectional = 2; None = 0. El objeto se cobra aunque GB = 0.\n"
        "4. Cross-cloud + ambos private: hourly $2.50, o $5.00 si Premium (hoja Regiones).\n"
        "5. Intra-cloud private: $0.50/h × multiplier de esa región.\n"
        "6. ClusterLinkingWrite = Kafka ingress del destination. No se cobra Kafka ingress extra.\n"
        "7. ClusterLinkingRead = $0.25/GB (oficial). Confluent paga el egress del cloud provider.\n"
        "8. Private-to-private Dedicated exige la MISMA Confluent Cloud Organization.\n"
        "9. Un cluster existente NO se convierte a Private Link. Jio no tiene PL.\n"
        "10. Schema Registry, CLI y provisioning de ksqlDB siguen necesitando Internet.\n"
        "11. Pega tus tarifas Dedicated reales en la hoja Tarifas.\n"
        "12. GB = GiB (2^30). Billing en UTC."
    )
    ws["A20"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A20"].font = Font(name="Calibri", size=10)
    ws["A20"].fill = fill(GRAY)

    # --- Consideraciones ---
    wc = wb.create_sheet("Consideraciones")
    wc["A1"] = "Consideraciones oficiales · Dedicated Private Link + Cluster Linking"
    wc["A1"].font = title_font
    wc.merge_cells("A1:C1")
    wc["A2"] = "Fuente"
    wc["B2"] = "Regla"
    wc["C2"] = "Detalle"
    style_header_bar(wc, 2, 1, 3)
    notes = [
        ("Billing Dimensions", "Hourly público", "$0.25/h por cluster con internet endpoints. Ambos clusters pagan si ambos son Confluent Cloud."),
        ("Billing Dimensions", "Hourly private intra-cloud", "$0.50/h por cluster privately networked. Bidirectional = 2 links al rate más alto ($0.50)."),
        ("Billing Dimensions", "Hourly cross-cloud private", "Aplica si (1) distinto cloud provider y (2) ambos usan private networking. $2.50/h por cluster, $5.00 si Premium region."),
        ("Billing Dimensions", "ClusterLinkingWrite", "Mismo precio que KafkaNetworkWrite del destination. No se cobra Kafka ingress adicional."),
        ("Billing Dimensions", "ClusterLinkingRead", "$0.25/GB incluyendo metadata. Confluent paga el egress del cloud provider. Hay volume discounts con account team."),
        ("Billing Dimensions", "Se cobra el objeto", "Hay cargo aunque no exista mirror topic ni throughput. Confluent no de-duplica links idénticos."),
        ("Billing Dimensions", "Regional multiplier", "Se aplica a precios base. Ejemplo docs: AWS us-east-1 = 1.0×, GCP asia-east2 = 1.2×. Los $2.50/$5.00 ya son el hourly cross-cloud; no se re-multiplica."),
        ("Billing Dimensions", "Storage RF=3", "Storage se factura post-replication. Dedicated replica 3 veces ≈ 3× el volumen escrito."),
        ("Billing Dimensions", "GiB y UTC", "1 GB = 2^30 bytes. Billing en UTC."),
        ("Cluster Linking private", "Misma organization", "Dedicated/Enterprise private → private (mismo o distinto cloud) exige la misma Confluent Cloud Organization. Cross-org no soportado."),
        ("Cluster Linking private", "Dedicated ↔ Dedicated private", "Soportado AWS↔Azure↔GCP. Source-initiated, destination-initiated o bidirectional."),
        ("Cluster Linking private", "Private → Public", "Requiere source-initiated link. No se crea desde Cloud Console en todos los casos."),
        ("Regions", "Premium region", "Fila Premium region (cross-cloud private networking) en las tablas de AWS/Azure/GCP. Solo encarece el hourly cross-cloud private."),
        ("Regions", "Jio", "jioindiawest / jioindiacentral: solo VNet Peering Dedicated. No hay Azure Private Link."),
        ("Cluster types", "CKU", "60 MBps ingress / 180 MBps egress / 4.500 partitions por CKU. Multi-zone mínimo 2 CKU. Single-zone 1+. Zone availability no se cambia después."),
        ("Cluster types", "Tope CKU", "Credit card: 4 CKU. Invoice: 24 CKU, más por request. Soft max AWS 252 / Azure 100 / GCP 152."),
        ("Azure PL Dedicated", "Endpoints", "Single-zone: 1 PE al Service Alias zonal. Multi-zone: 3 PE. Registrar subscription ID. El PE puede estar en otra región que el PLS."),
        ("Azure PL Dedicated", "DNS", "Private DNS recomendado. Zone name sin lkc- si el DNS domain lo trae. Records: * y *.az1/az2/az3. No hardcodear brokers."),
        ("Azure PL Dedicated", "No conversión", "Un cluster existente no se convierte a Private Link. DNS resolution tampoco se cambia después."),
        ("AWS PL Dedicated", "AZ matching", "Las AZ del VPC deben incluir las del cluster. Un VPC endpoint con subnet por AZ. Desmarcar Enable DNS name. SG: TCP 80, 443, 9092."),
        ("AWS PL Dedicated", "Sin PL cross-region", "Cross-region PrivateLink Attachment Connections no soportadas. use1-az3 en us-east-1 no se usa."),
        ("GCP PSC Dedicated", "Endpoints", "Single-zone: 1 PSC. Multi-zone: 3 PSC, uno por Service Attachment URI. Subnets en las mismas zonas."),
        ("GCP PSC Dedicated", "Attach same-region", "El attachment inicial no es cross-region. Después se puede activar PSC Global Access en el forwarding rule."),
        ("PL Dedicated (los 3)", "Internet residual", "Schema Registry, Confluent CLI (control plane) y provisioning de ksqlDB requieren outbound Internet. ksqlDB ya up habla por PL. Console necesita config extra."),
        ("PL Dedicated (los 3)", "Connectors", "Sources/sinks privados usan Egress Private Link Endpoints, no el inbound PL del cluster Kafka."),
        ("PL Dedicated (los 3)", "Flink / SR", "El inbound PL de Dedicated también puede usarse para algunos productos serverless (Flink, Schema Registry) según docs de cada servicio."),
    ]
    for i, (src, rule, detail) in enumerate(notes, 3):
        wc.cell(i, 1, src)
        wc.cell(i, 2, rule)
        wc.cell(i, 3, detail)
        wc.cell(i, 3).alignment = Alignment(wrap_text=True, vertical="top")
        for col in range(1, 4):
            wc.cell(i, col).border = thin
            wc.cell(i, col).font = Font(name="Calibri", size=10)
        wc.row_dimensions[i].height = 36
        if "Jio" in rule or "No conversión" in rule or "Misma organization" in rule:
            wc.cell(i, 2).fill = fill(AMBER)
    wc.column_dimensions["A"].width = 28
    wc.column_dimensions["B"].width = 28
    wc.column_dimensions["C"].width = 110
    wc.freeze_panes = "A3"
    wc.auto_filter.ref = f"A2:C{2+len(notes)}"
    wc.sheet_properties.tabColor = "F5A524"

    # Print area
    ws.print_area = "A1:L48"
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    out = "/Users/jackpaitan/github/ludhara-tech/confluent-cloud-clusterlink/calculadora-privatelink.xlsx"
    wb.save(out)
    return out


if __name__ == "__main__":
    path = build()
    print(path)
